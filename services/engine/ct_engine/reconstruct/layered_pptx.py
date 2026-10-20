"""Export a translated PDF as image backgrounds and editable text lines.

The PDF is the approved rendered output. Extracting its placed lines preserves
fit decisions. No scan pixels are claimed to be editable. Slide size is global
in PPTX, so mixed page sizes are rejected rather than silently distorted.
"""
from io import BytesIO
import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
from pptx.util import Pt


# Turned words cannot be a text box (see below), so they stay in the slide's
# background picture. The export used to refuse the whole document instead;
# now it says so on each slide where it happened.
TURNED_TEXT_NOTE = (
    "Some words on this slide are turned or printed at a slant. They are "
    "part of the background picture, so they show exactly as in the PDF but "
    "can't be edited here."
)
# A box she turned in the editor is turned text in the PDF, so it stays in
# the background picture like a tilted label. Making it a turned text box
# (python-pptx's shape.rotation) is not done: this export reads only the
# PDF, whose text gives an upright box and a direction, not the unturned
# frame, so the box could not be placed exactly. That would need the
# manifest's frames passed in -- a possible later step. Borders and fills
# are drawings, so they stay in the picture too, visible but not editable;
# inside space needs nothing, as lines are taken where they were drawn.
# Shapes likewise: outline and fill in the picture, their words text boxes.


# Underlines and highlights are painted shapes, which the background
# picture keeps; justified rows are drawn a word at a time, which comes back
# as one text box per word. Both are said on the slide, as for turned words.
STYLED_TEXT_NOTE = (
    "Underlines and highlights on this slide are part of the background "
    "picture, and justified lines are split into separate words. Edit them "
    "in Classroom Translate, or use the PDF."
)


def styled_pages(pages) -> set[int]:
    """Indexes of pages with an underline, a highlight or spread justified
    text, which this export cannot keep editable (see STYLED_TEXT_NOTE)."""
    from ..model.enums import TextAlignment
    from .fitting import justifiable

    out: set[int] = set()
    for page in pages:
        for t in page.text_objects:
            if t.drawn_layout is None or t.suppressed:
                continue
            spread = t.drawn_alignment is TextAlignment.JUSTIFY and justifiable(
                t.translated_text or ""
            )
            if t.manual_underline or t.manual_highlight or spread:
                out.add(page.index)
                break
    return out


def export_layered_pptx(pdf_bytes: bytes, style_notes: set[int] | None = None) -> bytes:
    """`style_notes`: indexes of pages to give STYLED_TEXT_NOTE."""
    presentation = Presentation()
    with fitz.open(stream=pdf_bytes, filetype="pdf") as pdf:
        if not len(pdf):
            raise ValueError("This document has no pages.")
        first = pdf[0].rect
        if any(abs(page.rect.width - first.width) > 0.1 or abs(page.rect.height - first.height) > 0.1 for page in pdf):
            raise ValueError("PowerPoint needs one page size. Download the PDF to keep this document's mixed page sizes.")
        if min(first.width, first.height) < 72 or max(first.width, first.height) > 4032:
            raise ValueError("This page size is outside PowerPoint's supported range. Use the PDF download.")
        presentation.slide_width, presentation.slide_height = Pt(first.width), Pt(first.height)
        for index, page in enumerate(pdf):
            blocks = page.get_text("dict")["blocks"]
            lines = [line for block in blocks if block.get("type") == 0 for line in block["lines"]]
            # Turned text stays in the background picture: visible exactly as
            # in the PDF, but not editable. A text box cannot be placed for
            # it -- the upright box get_text reports does not give back the
            # turned one -- and this export used to refuse the whole page
            # instead, which tilted labels in pictures (now drawn along
            # their slant) would have made common.
            level = [
                line for line in lines
                if abs(line["dir"][0] - 1) <= 0.01 and abs(line["dir"][1]) <= 0.01
            ]
            kept_in_picture = any(
                any(s["text"].strip() for s in line["spans"])
                for line in lines if line not in level
            )
            lines = level
            for line in lines:
                for span in line["spans"]:
                    page.add_redact_annot(span["bbox"], fill=None)
            page.apply_redactions(images=0, graphics=0, text=0)
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            png = BytesIO(page.get_pixmap(dpi=144, alpha=False).tobytes("png"))
            background = slide.shapes.add_picture(png, 0, 0, width=Pt(first.width), height=Pt(first.height))
            background.name = "Worksheet artwork (image)"
            # Said where she will see it: the slide she opens.
            notes = [
                *([TURNED_TEXT_NOTE] if kept_in_picture else []),
                *([STYLED_TEXT_NOTE] if style_notes and index in style_notes else []),
            ]
            if notes:
                slide.notes_slide.notes_text_frame.text = "\n\n".join(notes)
            for line_index, line in enumerate(lines):
                x0, y0, x1, y1 = line["bbox"]
                if not any(s["text"].strip() for s in line["spans"]):
                    continue
                box = slide.shapes.add_textbox(Pt(x0), Pt(y0), Pt(max(1, x1-x0) + 2), Pt(max(1, y1-y0) + 2))
                box.name = f"Page {index+1} text {line_index+1}"
                frame = box.text_frame
                frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
                frame.auto_size = MSO_AUTO_SIZE.NONE
                frame.word_wrap = False
                frame.vertical_anchor = MSO_ANCHOR.TOP
                paragraph = frame.paragraphs[0]
                paragraph.space_before = paragraph.space_after = Pt(0)
                for span in line["spans"]:
                    run = paragraph.add_run()
                    run.text = span["text"]
                    run.font.name = span["font"].split("+")[-1]
                    run.font.size = Pt(span["size"])
                    run.font.bold = bool(span["flags"] & 16)
                    run.font.italic = bool(span["flags"] & 2)
                    run.font.color.rgb = RGBColor.from_string(f'{span["color"] & 0xffffff:06X}')
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()
