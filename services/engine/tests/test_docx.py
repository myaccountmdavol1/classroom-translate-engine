"""DOCX ingest and reconstruction (phase 8).

The through-line of these tests is that a reflowable format fails
differently from a PDF. Nothing here can go wrong by drawing in the wrong
place; what can go wrong is text being skipped, duplicated, or written back
into the wrong run -- and all three are silent unless something checks.
"""

from __future__ import annotations

import hashlib
import html
import io
import re
import zipfile

import pytest
from conftest import corpus_file

from ct_engine.ingest.docx_parser import DocxParser
from ct_engine.ingest.docx_walk import (
    W,
    group_runs,
    iter_part_paragraphs,
    own_runs,
    parse_xml,
    split_paragraph,
)
from ct_engine.ingest.registry import parser_for
from ct_engine.ingest.validate import validate_upload
from ct_engine.model.enums import TranslationStatus
from ct_engine.reconstruct.docx_builder import rebuild_docx

WORKSHEET = "20_worksheet.docx"
TWO_SECTIONS = "21_two_sections.docx"


def _open(name: str):
    return DocxParser().open_bytes(corpus_file(name))


def _pages(name: str):
    with _open(name) as doc:
        return [doc.parse_page(i) for i in range(doc.page_count)]


def _all_text(name: str) -> list[str]:
    return [o.text for p in _pages(name) for o in p.text_objects]


def _translate_all(pages, fn=lambda t: f"<{t}>"):
    for page in pages:
        for obj in page.text_objects:
            obj.translated_text = fn(obj.text)
            obj.translation_status = TranslationStatus.TRANSLATED
    return pages


def _w_t_values(data: bytes, part: str = "word/document.xml") -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        xml = z.read(part).decode()
    # Unescape: the marker characters these tests use are entities in the XML.
    return [html.unescape(t) for t in re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml)]


def _body_xml(data: bytes, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return html.unescape(z.read(part).decode())


# --- The walk -------------------------------------------------------------


def test_hyperlink_text_is_seen_exactly_once():
    """python-docx puts link text in `.text` but not in `.runs`.

    Translating `.text` and writing it over `.runs` would emit the sentence
    twice, half of it untranslated. The walk has to report link runs.
    """
    texts = _all_text(WORKSHEET)
    joined = " ".join(texts)
    assert joined.count("the lab safety guide") == 1
    # ...and as its own block, so emptying a neighbour cannot orphan the link.
    assert "the lab safety guide" in texts


def test_text_box_content_is_found():
    """`Document.paragraphs` reports nothing for a text box."""
    assert any("label every axis" in t for t in _all_text(WORKSHEET))


def test_headers_footers_and_table_cells_are_found():
    texts = _all_text(WORKSHEET)
    assert any("Ms. Alvarez" in t for t in texts), "header missing"
    assert any("Riverside Middle School" in t for t in texts), "footer missing"
    assert any(t.strip() == "Mass before (g)" for t in texts), "table cell missing"


def test_invisible_run_splits_do_not_look_like_formatting():
    """Word splits runs with no visible effect; those must merge.

    Without this the 'Directions:' paragraph would look like three different
    formats and lose its bold lead-in for nothing.
    """
    with _open(WORKSHEET) as doc:
        root = doc._body
    refs = [r for r in iter_part_paragraphs(root, "word/document.xml")]
    directions = next(r for r in refs if r.units and
                      r.units[0].text.startswith("Directions:"))
    assert len(directions.units) == 2, "expected label + body"
    assert directions.units[0].text.strip() == "Directions:"
    assert not any(u.collapsed for u in directions.units)


def test_interleaved_bold_collapses_and_says_so():
    """A bold word mid-sentence cannot survive, and must not pretend to."""
    page = _pages(WORKSHEET)[0]
    obj = next(o for o in page.text_objects if "mitochondria" in o.text)
    assert obj.id in page.diagnostics.extra["docx_collapsed_formatting"]


def test_a_section_is_a_page():
    with _open(TWO_SECTIONS) as doc:
        assert doc.page_count == 2
    pages = _pages(TWO_SECTIONS)
    assert pages[1].width > pages[1].height, "second section is landscape"


def test_pages_are_marked_as_having_no_layout():
    """Fitting, reflow and collision QA must never run on these boxes."""
    assert all(not p.has_fixed_layout for p in _pages(WORKSHEET))


def test_no_page_previews_are_claimed():
    with _open(WORKSHEET) as doc:
        assert doc.supports_preview is False


# --- Write-back -----------------------------------------------------------


def test_round_trip_translates_every_block_exactly_once():
    src = corpus_file(WORKSHEET)
    pages = _translate_all(_pages(WORKSHEET))
    result = rebuild_docx(src, pages)
    assert result.unresolved == []

    out_text = " ".join(_w_t_values(result.data))
    # Every translated block appears, and no source block survives beside it.
    assert out_text.count("<1. What is the main function of the mitochondria "
                          "in a plant cell?>") == 1
    assert "the lab safety guide" not in out_text.replace(
        "<the lab safety guide>", ""
    )


def test_paragraph_count_is_preserved():
    """The complete-document guarantee for a format with no page count."""
    src = corpus_file(WORKSHEET)
    with DocxParser().open_bytes(src) as doc:
        before = doc.paragraph_count
    result = rebuild_docx(src, _translate_all(_pages(WORKSHEET)))
    with DocxParser().open_bytes(result.data) as doc:
        assert doc.paragraph_count == before


def test_images_are_byte_identical():
    """Never regenerate artwork (spec sec. 14). Nothing is even extracted."""
    src = corpus_file(WORKSHEET)
    result = rebuild_docx(src, _translate_all(_pages(WORKSHEET)))

    def media(data: bytes) -> dict[str, str]:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return {
                n: hashlib.sha256(z.read(n)).hexdigest()
                for n in z.namelist()
                if n.startswith("word/media/")
            }

    before = media(src)
    assert before, "fixture has no embedded image"
    assert media(result.data) == before


def test_hyperlink_target_survives():
    src = corpus_file(WORKSHEET)
    result = rebuild_docx(src, _translate_all(_pages(WORKSHEET)))
    with zipfile.ZipFile(io.BytesIO(result.data)) as z:
        rels = z.read("word/_rels/document.xml.rels").decode()
    body = _body_xml(result.data)
    assert "https://example.org/lab-safety" in rels
    assert "<w:hyperlink" in body
    assert "<the lab safety guide>" in body


def test_tabs_are_re_emitted_as_tab_elements():
    """`Name: ___\\tDate: ___` is two fields; a space would merge them."""
    src = corpus_file(WORKSHEET)
    pages = _pages(WORKSHEET)
    obj = next(o for p in pages for o in p.text_objects if "Date:" in o.text)
    _translate_all(pages, lambda t: t)  # identity: keeps the tab
    result = rebuild_docx(src, pages)
    body = _body_xml(result.data)
    assert "\t" in obj.text, "the source tab must reach the translator"
    assert "<w:tab/>" in body


def test_leading_space_is_preserved():
    """Word strips edge spaces without xml:space, joining label to body."""
    src = corpus_file(WORKSHEET)
    pages = _translate_all(_pages(WORKSHEET), lambda t: t)
    result = rebuild_docx(src, pages)
    assert 'xml:space="preserve"' in _body_xml(result.data)


def test_protected_blocks_are_left_untouched():
    src = corpus_file(WORKSHEET)
    pages = _pages(WORKSHEET)
    for page in pages:
        for obj in page.text_objects:
            obj.translation_status = TranslationStatus.PROTECTED
            obj.translated_text = "SHOULD NOT APPEAR"
    result = rebuild_docx(src, pages)
    assert "SHOULD NOT APPEAR" not in " ".join(_w_t_values(result.data))
    assert "Cells and Energy: Review Packet" in " ".join(_w_t_values(result.data))


def test_suppressed_block_is_removed_not_left_behind():
    src = corpus_file(WORKSHEET)
    pages = _translate_all(_pages(WORKSHEET), lambda t: t)
    target = next(o for p in pages for o in p.text_objects
                  if "Do not write below" in o.text)
    target.suppressed = True
    result = rebuild_docx(src, pages)
    assert "Do not write below" not in " ".join(_w_t_values(result.data))


def test_manual_colour_and_size_reach_the_document():
    src = corpus_file(WORKSHEET)
    pages = _translate_all(_pages(WORKSHEET), lambda t: t)
    # A body block, not the header: the header lives in its own part.
    target = next(o for o in pages[0].text_objects
                  if o.anchor.startswith("word/document.xml"))
    target.manual_color = "#1155cc"
    target.manual_font_size = 19.0
    result = rebuild_docx(src, pages)
    body = _body_xml(result.data)
    assert 'w:val="1155CC"' in body
    assert 'w:val="38"' in body  # half-points


def test_output_is_a_readable_word_file():
    """Reopen with a different library than the one that wrote it."""
    pytest.importorskip("docx")
    from docx import Document

    src = corpus_file(WORKSHEET)
    result = rebuild_docx(src, _translate_all(_pages(WORKSHEET)))
    doc = Document(io.BytesIO(result.data))
    assert any("<" in p.text for p in doc.paragraphs)
    assert doc.tables


def test_unresolvable_anchor_is_reported_not_swallowed():
    src = corpus_file(WORKSHEET)
    pages = _translate_all(_pages(WORKSHEET))
    pages[0].text_objects[0].anchor = "word/document.xml#9999#0"
    result = rebuild_docx(src, pages)
    assert result.unresolved == ["word/document.xml#9999#0"]
    assert any(f.code == "docx_anchor_lost" for f in result.stats[0].findings)


# --- The door -------------------------------------------------------------


def test_upload_validation_accepts_docx():
    r = validate_upload(corpus_file(WORKSHEET), "packet.docx")
    assert r.ok
    assert r.content_type.endswith("wordprocessingml.document")


def test_registry_picks_the_docx_parser():
    assert parser_for(None, "a.docx").name == "docx"
    assert parser_for(None, "a.pdf").name != "docx"


# --- Through the real pipeline --------------------------------------------


@pytest.mark.slow
def test_full_word_translation_round_trip(client):
    """Upload -> analyze -> translate -> write back -> download, for real.

    The assertions that matter here are the ones a unit test cannot make:
    that the job reaches COMPLETE (so the completion gate passed), that the
    download is served as a Word file rather than a PDF, and that the file a
    teacher receives opens and has every paragraph the original had.
    """
    import time

    resp = client.post(
        "/api/documents",
        files={
            "file": (
                WORKSHEET,
                corpus_file(WORKSHEET),
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 200, resp.text
    doc_id = resp.json()["document_id"]

    deadline = time.time() + 120
    while time.time() < deadline:
        job = client.get(f"/api/documents/{doc_id}/job").json()
        if job["stage"] in ("REVIEW_READY", "COMPLETE", "FAILED"):
            break
        time.sleep(0.05)
    assert job["stage"] != "FAILED", job

    started = client.post(
        f"/api/documents/{doc_id}/translate", json={"target_language": "es"}
    )
    assert started.status_code == 200, started.text

    deadline = time.time() + 180
    while time.time() < deadline:
        job = client.get(f"/api/documents/{doc_id}/job").json()
        if job["stage"] in ("COMPLETE", "FAILED"):
            break
        time.sleep(0.05)
    assert job["stage"] == "COMPLETE", job

    dl = client.get(f"/api/documents/{doc_id}/download")
    assert dl.status_code == 200
    assert dl.headers["content-type"].endswith("wordprocessingml.document")
    assert dl.headers["content-disposition"].endswith('.docx"')
    assert "Spanish" in dl.headers["content-disposition"]

    with DocxParser().open_bytes(corpus_file(WORKSHEET)) as before:
        expected = before.paragraph_count
    with DocxParser().open_bytes(dl.content) as after:
        assert after.paragraph_count == expected

    text = " ".join(_w_t_values(dl.content))
    # Something was actually translated...
    assert "Instrucciones" in text or "Direcciones" in text, text[:400]
    # ...the equation was not touched...
    assert "3x + 7 = 22" in text
    # ...the answer blanks survived...
    assert "______" in text
    # ...and the artwork is the original bytes.
    with zipfile.ZipFile(io.BytesIO(dl.content)) as z:
        out_media = {n: hashlib.sha256(z.read(n)).hexdigest()
                     for n in z.namelist() if n.startswith("word/media/")}
    with zipfile.ZipFile(io.BytesIO(corpus_file(WORKSHEET))) as z:
        src_media = {n: hashlib.sha256(z.read(n)).hexdigest()
                     for n in z.namelist() if n.startswith("word/media/")}
    assert out_media == src_media


@pytest.mark.slow
def test_editing_a_word_document_reaches_the_download(client):
    """A correction must survive into the file the teacher downloads.

    This is the path that broke first: the editor re-rendered the page with
    the PDF engine to refresh a preview, which a Word file does not have, so
    every correction was rejected with "that page could not be rebuilt".
    """
    import time

    resp = client.post(
        "/api/documents",
        files={
            "file": (
                TWO_SECTIONS,
                corpus_file(TWO_SECTIONS),
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
    )
    doc_id = resp.json()["document_id"]
    deadline = time.time() + 120
    while time.time() < deadline:
        job = client.get(f"/api/documents/{doc_id}/job").json()
        if job["stage"] in ("REVIEW_READY", "COMPLETE", "FAILED"):
            break
        time.sleep(0.05)
    client.post(f"/api/documents/{doc_id}/translate", json={"target_language": "es"})
    deadline = time.time() + 180
    while time.time() < deadline:
        job = client.get(f"/api/documents/{doc_id}/job").json()
        if job["stage"] in ("COMPLETE", "FAILED"):
            break
        time.sleep(0.05)
    assert job["stage"] == "COMPLETE", job

    blocks = client.get(f"/api/documents/{doc_id}/pages/1/blocks").json()["blocks"]
    target = next(b for b in blocks if "Stomata" in b["original"])
    corrected = "Estomas. Aberturas en una hoja que dejan entrar y salir los gases."

    patched = client.patch(
        f"/api/documents/{doc_id}/pages/1/blocks/{target['id']}",
        json={"text": corrected},
    )
    assert patched.status_code == 200, patched.text

    dl = client.get(f"/api/documents/{doc_id}/download")
    assert dl.status_code == 200
    assert corrected in " ".join(_w_t_values(dl.content))


def test_an_external_entity_cannot_read_a_file_off_the_server():
    """A crafted upload must not be able to exfiltrate local files.

    These packages are XML supplied by the public, and the classic attack is
    to declare an entity pointing at a local path so its contents come back
    inside a "translation" the attacker then downloads.

    The guarantee asserted here is the one that matters: the file is never
    read, so its contents appear nowhere -- not in a block, not in the
    rebuilt document, not in an error message. The document is NOT required
    to be rejected. lxml with `resolve_entities=False` leaves a declared
    entity as an unresolved reference rather than failing, which is a
    perfectly safe outcome, and a test that demanded a rejection would be
    asserting an implementation detail instead of the property.
    """
    import pathlib
    import tempfile

    with tempfile.NamedTemporaryFile("w", suffix=".txt", delete=False) as fh:
        fh.write("SUPER-SECRET-VALUE")
        secret = pathlib.Path(fh.name)
    try:
        src = corpus_file(WORKSHEET)
        out = io.BytesIO()
        built = False
        with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(out, "w") as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    body = data.decode()
                    doctype = (
                        "<!DOCTYPE w:document [ "
                        f'<!ENTITY xxe SYSTEM "file://{secret}"> ]>'
                    )
                    head, sep, rest = body.partition("?>")
                    assert sep, "fixture has no XML declaration"
                    body = f"{head}{sep}\n{doctype}{rest}"
                    body = body.replace(
                        "Cells and Energy: Review Packet", "&xxe;", 1
                    )
                    assert "&xxe;" in body and doctype in body
                    built = True
                    data = body.encode()
                zout.writestr(item, data)
        assert built, "the hostile payload was never constructed"
        payload = out.getvalue()

        result = validate_upload(payload, "attack.docx")
        assert "SUPER-SECRET-VALUE" not in (result.message or "")
        assert "SUPER-SECRET-VALUE" not in (result.diagnostic or "")

        try:
            with DocxParser().open_bytes(payload) as doc:
                pages = [doc.parse_page(i) for i in range(doc.page_count)]
            text = " ".join(o.text for p in pages for o in p.text_objects)
            rebuilt = rebuild_docx(payload, _translate_all(pages)).data
        except Exception as exc:
            text, rebuilt = str(exc), b""
        assert "SUPER-SECRET-VALUE" not in text
        assert b"SUPER-SECRET-VALUE" not in rebuilt
    finally:
        secret.unlink(missing_ok=True)
