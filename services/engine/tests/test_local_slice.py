"""Synthetic checks for images, portable projects and editable export."""
import io
import json
import zipfile
import fitz
import pytest
from PIL import Image
from pptx import Presentation
from ct_engine.ingest.image_import import image_to_pdf
from ct_engine.pipeline.project_package import export_project, import_project
from ct_engine.pipeline.manifest_store import ManifestStore
from ct_engine.model.document import DocumentManifest
from ct_engine.reconstruct.layered_pptx import export_layered_pptx
from ct_engine.util.hashing import sha256_bytes


def test_image_import_preserves_dimensions():
    raw = io.BytesIO()
    Image.new("RGB", (600, 800), "white").save(raw, format="PNG", dpi=(150, 150))
    pdf, mime = image_to_pdf(raw.getvalue())
    assert mime == "image/png"
    with fitz.open(stream=pdf, filetype="pdf") as document:
        assert len(document) == 1
        assert document[0].rect.width == pytest.approx(288, abs=0.1)
        image = document.extract_image(document[0].get_images()[0][0])
        assert (image["width"], image["height"]) == (600, 800)
    with pytest.raises(ValueError):
        image_to_pdf(b"broken")


def test_project_roundtrip_and_tampering(store):
    old = "a" * 32
    source = f"documents/{old}/source.pdf"
    with fitz.open() as pdf:
        pdf.new_page().insert_text((50, 50), "Synthetic worksheet")
        data = pdf.tobytes()
    store.put_bytes(source, data)
    manifest = DocumentManifest(document_id=old, filename="synthetic.pdf", content_type="application/pdf",
        byte_size=len(data), sha256=sha256_bytes(data), source_key=source)
    manifest.analysis.page_count = 1
    ManifestStore(store).save_manifest(manifest)
    package = export_project(old, store)
    reopened = import_project(package, store)
    assert reopened.document_id != old
    assert store.get_bytes(reopened.source_key) == data
    assert store.get_bytes(source) == data
    damaged = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(package)) as original, zipfile.ZipFile(damaged, "w") as altered:
        for name in original.namelist():
            altered.writestr(name, b"tampered" if name == "source.pdf" else original.read(name))
    with pytest.raises(ValueError, match="checksum"):
        import_project(damaged.getvalue(), store)


def test_project_rejects_traversal_and_unknown_version(store):
    for header, path in [({"schema_version": 99}, "safe"), ({"schema_version": 1}, "../escape")]:
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w") as archive:
            archive.writestr("project.json", json.dumps(header))
            archive.writestr(path, "content")
        with pytest.raises(ValueError):
            import_project(data.getvalue(), store)


def test_pptx_has_editable_text_and_background_without_duplicate_letters():
    with fitz.open() as pdf:
        page = pdf.new_page(width=612, height=792)
        page.insert_text((72, 90), "Texto editable", fontsize=16)
        page.draw_line((72, 130), (500, 130))
        data = pdf.tobytes()
    deck = Presentation(io.BytesIO(export_layered_pptx(data)))
    assert len(deck.slides) == 1
    assert deck.slide_width == 612 * 12700
    shapes = deck.slides[0].shapes
    assert [s.text for s in shapes if s.has_text_frame] == ["Texto editable"]
    background = Image.open(io.BytesIO(shapes[0].image.blob)).convert("RGB")
    assert background.crop((140, 140, 400, 185)).getextrema()[0][0] == 255
    assert background.crop((140, 258, 1000, 262)).getextrema()[0][0] < 128


def test_pptx_refuses_to_distort_mixed_sizes():
    with fitz.open() as pdf:
        pdf.new_page(width=612, height=792)
        pdf.new_page(width=792, height=612)
        data = pdf.tobytes()
    with pytest.raises(ValueError, match="one page size"):
        export_layered_pptx(data)



def test_transparent_image_is_composited_on_white():
    raw = io.BytesIO()
    Image.new("RGBA", (100, 100), (0, 0, 0, 0)).save(raw, format="PNG")
    pdf, _ = image_to_pdf(raw.getvalue())
    with fitz.open(stream=pdf, filetype="pdf") as document:
        pixels = document[0].get_pixmap().samples
        assert min(pixels) == 255
