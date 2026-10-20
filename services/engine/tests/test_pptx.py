"""PPTX ingest and reconstruction (phase 8).

A deck fails differently again. Its coordinates are real, so the review
overlay has to be right; and PowerPoint will not reflow, so a translation
that outgrows its box is a defect the engine has to notice rather than a
cosmetic difference.
"""

from __future__ import annotations

import hashlib
import html
import io
import re
import zipfile

import pytest
from conftest import corpus_file

from ct_engine.ingest.pptx_parser import PptxParser
from ct_engine.ingest.registry import parser_for
from ct_engine.ingest.validate import validate_upload
from ct_engine.model.enums import TranslationStatus
from ct_engine.reconstruct.pptx_builder import MIN_FONT_SCALE, rebuild_pptx

DECK = "22_lesson.pptx"


def _pages():
    with PptxParser().open_bytes(corpus_file(DECK)) as doc:
        return [doc.parse_page(i) for i in range(doc.page_count)]


def _texts(pages) -> list[str]:
    return [o.text for p in pages for o in p.text_objects]


def _translate(pages, fn):
    for page in pages:
        for obj in page.text_objects:
            obj.translated_text = fn(obj.text)
            obj.translation_status = TranslationStatus.TRANSLATED
    return pages


def _slide_xml(data: bytes, n: int = 1) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return z.read(f"ppt/slides/slide{n}.xml").decode()


def _all_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        parts = [n for n in z.namelist() if n.startswith("ppt/slides/slide")]
        return " ".join(
            html.unescape(t)
            for part in parts
            for t in re.findall(r"<a:t>([^<]*)</a:t>", z.read(part).decode())
        )


# --- Parsing --------------------------------------------------------------


def test_one_slide_is_one_page():
    pages = _pages()
    assert len(pages) == 3
    assert [p.number for p in pages] == [1, 2, 3]


def test_slides_have_real_coordinates():
    """Unlike a Word section: a shape's box is honoured exactly."""
    assert all(p.has_fixed_layout for p in _pages())


def test_placeholder_geometry_is_inherited_from_the_master():
    """Most text in a deck has no box of its own.

    PowerPoint stores it on the layout or master, so a parser that reads
    only the slide would find no position for the title or the body -- which
    is nearly all of it.
    """
    title = next(
        o for o in _pages()[0].text_objects if "Photosynthesis" in o.text
    )
    assert title.bbox.width > 100, "title has no resolved box"
    assert title.style.font_size > 30, "title kept the body default size"


def test_table_cells_form_a_grid_not_a_column():
    cells = [o for o in _pages()[2].text_objects if o.role.value == "TABLE_CELL"]
    assert len(cells) == 9
    tops = sorted({round(c.bbox.y0) for c in cells})
    lefts = sorted({round(c.bbox.x0) for c in cells})
    assert len(tops) == 3 and len(lefts) == 3, (tops, lefts)


def test_no_two_blocks_claim_the_same_box():
    """Identical boxes read to QA as text printing over text."""
    for page in _pages():
        seen = set()
        for o in page.text_objects:
            key = (round(o.bbox.x0), round(o.bbox.y0), round(o.bbox.x1), round(o.bbox.y1))
            assert key not in seen, f"duplicate box on page {page.number}: {o.text!r}"
            seen.add(key)


def test_hyperlink_text_is_its_own_block():
    texts = _texts(_pages())
    assert "the study guide" in texts


def test_speaker_notes_are_translated_but_not_placed_on_the_slide():
    page = _pages()[0]
    note = next(o for o in page.text_objects if "Remind students" in o.text)
    # Below every shape that is actually drawn on the slide.
    drawn = [o for o in page.text_objects if o is not note]
    assert note.bbox.y0 >= max(o.bbox.y1 for o in drawn)


def test_no_slide_previews_are_claimed():
    with PptxParser().open_bytes(corpus_file(DECK)) as doc:
        assert doc.supports_preview is False


# --- Write-back -----------------------------------------------------------


def test_round_trip_keeps_every_paragraph():
    src = corpus_file(DECK)
    with PptxParser().open_bytes(src) as doc:
        before = doc.paragraph_count
    result = rebuild_pptx(src, _translate(_pages(), lambda t: f"<{t}>"))
    assert result.unresolved == []
    with PptxParser().open_bytes(result.data) as doc:
        assert doc.paragraph_count == before
        assert doc.page_count == 3


def test_images_are_byte_identical():
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: f"<{t}>"))

    def media(data: bytes) -> dict[str, str]:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return {
                n: hashlib.sha256(z.read(n)).hexdigest()
                for n in z.namelist()
                if n.startswith("ppt/media/")
            }

    assert media(src), "fixture has no picture"
    assert media(result.data) == media(src)


def test_hyperlink_target_survives():
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: f"<{t}>"))
    with zipfile.ZipFile(io.BytesIO(result.data)) as z:
        rels = z.read("ppt/slides/_rels/slide2.xml.rels").decode()
    assert "https://example.org/study-guide" in rels
    assert "<the study guide>" in _all_text(result.data)


def test_speaker_notes_reach_the_notes_part():
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: f"<{t}>"))
    with zipfile.ZipFile(io.BytesIO(result.data)) as z:
        notes = [n for n in z.namelist() if "notesSlide" in n and n.endswith(".xml")]
        assert notes
        body = " ".join(z.read(n).decode() for n in notes)
    assert "Remind students" in body and "&lt;" in body or "<Remind" in body


def test_text_that_still_fits_is_left_alone():
    """Never shrink what does not need shrinking."""
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: t))
    assert sum(s.shrunk for s in result.stats.values()) == 0
    assert "<a:normAutofit" not in _slide_xml(result.data, 2)


def test_text_that_outgrows_its_box_is_shrunk_the_way_powerpoint_would():
    """PowerPoint does not reflow, so a longer translation must be fitted.

    The shrink is written as `a:normAutofit`, the same element PowerPoint
    writes for "Shrink text on overflow", so the deck renders the same way
    everywhere rather than relying on this engine's own layout guess.
    """
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: (t + " ") * 4))
    assert sum(s.shrunk for s in result.stats.values()) > 0
    assert "<a:normAutofit" in _slide_xml(result.data, 1)


def test_a_shape_set_to_grow_with_its_text_is_not_touched():
    """The author already said what should happen; do not override it."""
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: (t + " ") * 6))
    xml = _slide_xml(result.data, 2)
    assert "<a:spAutoFit/>" in xml, "fixture lost its grow-to-fit box"


def test_text_too_long_even_when_shrunk_is_flagged_not_squashed():
    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: (t + " ") * 25))
    findings = [f for s in result.stats.values() for f in s.findings]
    assert any(f.code == "pptx_text_may_not_fit" for f in findings)
    scales = [
        int(m) for m in re.findall(r'fontScale="(\d+)"', _slide_xml(result.data, 1))
    ]
    assert scales and min(scales) >= MIN_FONT_SCALE * 100_000 - 1, scales


def test_protected_blocks_are_left_untouched():
    src = corpus_file(DECK)
    pages = _pages()
    for page in pages:
        for obj in page.text_objects:
            obj.translation_status = TranslationStatus.PROTECTED
            obj.translated_text = "SHOULD NOT APPEAR"
    result = rebuild_pptx(src, pages)
    text = _all_text(result.data)
    assert "SHOULD NOT APPEAR" not in text
    assert "Photosynthesis: How Plants Make Food" in text


def test_manual_colour_and_size_reach_the_slide():
    src = corpus_file(DECK)
    pages = _translate(_pages(), lambda t: t)
    target = next(o for o in pages[0].text_objects if "Photosynthesis" in o.text)
    target.manual_color = "#1155cc"
    target.manual_font_size = 28.0
    result = rebuild_pptx(src, pages)
    xml = _slide_xml(result.data, 1)
    assert 'val="1155CC"' in xml
    assert 'sz="2800"' in xml


def test_output_is_a_readable_presentation():
    pytest.importorskip("pptx")
    from pptx import Presentation

    src = corpus_file(DECK)
    result = rebuild_pptx(src, _translate(_pages(), lambda t: f"<{t}>"))
    prs = Presentation(io.BytesIO(result.data))
    assert len(prs.slides) == 3
    assert any(
        "<" in sh.text_frame.text
        for sl in prs.slides
        for sh in sl.shapes
        if sh.has_text_frame
    )


def test_unresolvable_anchor_is_reported_not_swallowed():
    src = corpus_file(DECK)
    pages = _translate(_pages(), lambda t: t)
    pages[0].text_objects[0].anchor = "ppt/slides/slide1.xml#9999#0"
    result = rebuild_pptx(src, pages)
    assert result.unresolved == ["ppt/slides/slide1.xml#9999#0"]


# --- The door -------------------------------------------------------------


def test_upload_validation_accepts_pptx():
    r = validate_upload(corpus_file(DECK), "lesson.pptx")
    assert r.ok and r.page_count == 3
    assert r.content_type.endswith("presentationml.presentation")


def test_a_deck_named_docx_is_refused_clearly():
    r = validate_upload(corpus_file(DECK), "lesson.docx")
    assert not r.ok and "PowerPoint" in r.message


def test_registry_picks_the_pptx_parser():
    assert parser_for(None, "a.pptx").name == "pptx"
