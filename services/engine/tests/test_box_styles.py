"""Google Slides' box styles: a border around the box, space inside it,
and turning any box to any angle.

Each is a choice she makes in the editor. The engine draws it in the PDF
and in the block's live picture alike, keeps it through undo and Reset, and
carries it into the Office pictures; Word and PowerPoint's own text keeps
its own box, on purpose.
"""

from __future__ import annotations

import io
import math
from types import SimpleNamespace

import fitz
import numpy as np
import pytest
from PIL import Image

from ct_engine.model.document import DocumentManifest
from ct_engine.model.enums import TextSource, TranslationStatus
from ct_engine.model.geometry import Rect
from ct_engine.model.objects import DrawnLayout, TextObject, grow_quad
from ct_engine.pipeline.edit import (
    EditError,
    _angle,
    _clamped,
    _snapshot,
    add_block,
    apply_bulk,
    apply_edit,
    apply_layout,
    apply_positions,
    redo,
    reset_edit,
    undo,
)
from ct_engine.pipeline.reconstruct_stage import reconstruct_and_export
from ct_engine.reconstruct.builder import _turn_point, frame_corners, turned_frame_of
from ct_engine.reconstruct.fitting import MIN_INNER, inset
from ct_engine.reconstruct.sprite import render_sprite
from ct_engine.util.hashing import sha256_bytes
from test_edit import prepared
from test_text_styles import BOX, SPANISH, _directions, _page_pixels, _styled
from test_tilted_text import _read, _translate, sprite_on_page

BOX_FIELDS = (
    "manual_angle", "manual_padding", "manual_border_color",
    "manual_border_width", "manual_border_dash",
)
DPI = 144  # 2 pixels a point


def _px(v):
    return int(round(v * DPI / 72))


def _red(arr):
    return (arr[..., 0] > 190) & (arr[..., 1] < 90) & (arr[..., 2] < 90)


def _blue(arr):
    return (arr[..., 2] > 190) & (arr[..., 0] < 90) & (arr[..., 1] < 90)


def _runs(row):
    """Lengths of the True runs and of the False gaps between them."""
    edges = np.diff(np.concatenate(([0], row.astype(np.int8), [0])))
    starts, ends = np.where(edges == 1)[0], np.where(edges == -1)[0]
    on = ends - starts
    off = starts[1:] - ends[:-1]
    return on, off


def _render(manifest, store):
    plates: dict = {}
    final, out = _page_pixels(manifest, store, dpi=DPI, plates=plates)
    return final, plates[0], out


def _label_manifest(store, angle=13.0, text="ESCUELA HARBOR", **kw):
    """A picture label read from a tilted_scenes picture, translated, in a
    manifest the edit functions can rebuild."""
    pdf, page, _, _, (label,) = _read(angle, **kw)
    _translate(label, text)
    store.put_bytes("documents/tilt/source.pdf", pdf)
    manifest = DocumentManifest(
        document_id="tilt", filename="t.pdf", content_type="application/pdf",
        byte_size=len(pdf), sha256=sha256_bytes(pdf),
        source_key="documents/tilt/source.pdf", pages=[page],
    )
    manifest.analysis.page_count = 1
    reconstruct_and_export(manifest, store=store, target_lang="es")
    return manifest, label


def _edit(manifest, obj, store, **kw):
    return apply_layout(manifest, 1, obj.id, store=store, target_lang="es", **kw)


def _aabb(points):
    xs, ys = zip(*points)
    return Rect(x0=min(xs), y0=min(ys), x1=max(xs), y1=max(ys))


def _close(a, b, tol=0.01):
    return all(abs(x - y) <= tol for x, y in zip(a.as_tuple(), b.as_tuple()))


# --- 1. Border --------------------------------------------------------------
def test_a_solid_border_is_drawn_along_the_frame_at_its_weight(store):
    manifest, obj = _styled(store, border_color="#FF0000", border_width=4)
    assert obj.manual_border_color == "#ff0000" and obj.manual_border_width == 4.0
    frame = obj.drawn_bbox
    assert _close(frame, Rect(**BOX))
    final, plate, _ = _render(manifest, store)
    red = _red(final)
    # Across the top edge at the middle: about 4pt = 8 pixels, centred on it.
    column = red[_px(frame.y0) - 8:_px(frame.y0) + 8, _px((frame.x0 + frame.x1) / 2)]
    assert 6 <= column.sum() <= 10
    for x, y in ((frame.x0, (frame.y0 + frame.y1) / 2), (frame.x1, frame.y1 - 10)):
        assert red[_px(y), _px(x)]
    # Nothing inside the frame but the words, which are not red.
    inner = frame.expanded(-(2 + 1))
    assert not red[_px(inner.y0):_px(inner.y1), _px(inner.x0):_px(inner.x1)].any()
    # The editor's clean background never holds a border.
    assert not _red(plate).any()


def test_dashes_are_three_weights_long_with_gaps_of_two(store):
    manifest, obj = _styled(store, border_color="#ff0000", border_width=4, border_dash="dashed")
    assert obj.manual_border_dash == "dashed"
    frame = obj.drawn_bbox
    final, _, _ = _render(manifest, store)
    row = _red(final)[_px(frame.y0), _px(frame.x0) + 8:_px(frame.x1) - 8]
    on, off = _runs(row)
    # Inner runs only: the first and last may be cut by the window.
    assert len(on) >= 4
    assert np.all(np.abs(on[1:-1] - 24) <= 3)
    assert np.all(np.abs(off - 16) <= 3)


def test_dots_are_round_and_two_weights_apart(store):
    manifest, obj = _styled(store, border_color="#ff0000", border_width=4, border_dash="dotted")
    frame = obj.drawn_bbox
    final, _, _ = _render(manifest, store)
    red = _red(final)
    y = _px(frame.y0)
    row = red[y, _px(frame.x0) + 8:_px(frame.x1) - 8]
    on, off = _runs(row)
    assert len(on) >= 8
    assert np.all(np.abs(on[1:-1] - 8) <= 3) and np.all(np.abs(off - 8) <= 3)
    # Round, not square: a dot fills about pi/4 of the square around it.
    start = _px(frame.x0) + 8 + int(np.where(np.diff(row.astype(int)) == 1)[0][1]) + 1
    square = red[y - 4:y + 4, start:start + 8]
    assert 0.6 < square.mean() < 0.92


def test_the_border_is_on_top_of_her_background(store):
    manifest, obj = _styled(store, border_color="#ff0000", border_width=4, fill="#ffff00")
    frame = obj.drawn_bbox
    final, _, _ = _render(manifest, store)
    mid = (frame.x0 + frame.x1) / 2
    assert _red(final)[_px(frame.y0), _px(mid)]
    inside = final[_px(frame.y0 + 4), _px(mid)]
    assert inside[0] > 200 and inside[1] > 200 and inside[2] < 80  # yellow


def test_a_turned_box_has_its_border_on_the_turned_frame(store):
    manifest, obj = _styled(store, border_color="#ff0000", border_width=3)
    _edit(manifest, obj, store, angle=30)
    quad = obj.drawn_layout.quad
    assert quad and obj.drawn_layout.border_width == 3.0
    final, _, _ = _render(manifest, store)
    red = _red(final)
    for (ax, ay), (bx, by) in zip(quad, quad[1:] + quad[:1]):
        mx, my = (ax + bx) / 2, (ay + by) / 2
        assert red[_px(my) - 1:_px(my) + 2, _px(mx) - 1:_px(mx) + 2].any()
    box = obj.drawn_bbox
    for x, y in ((box.x0 + 2, box.y0 + 2), (box.x1 - 2, box.y1 - 2)):
        assert not red[_px(y), _px(x)]


@pytest.mark.parametrize("angle", [None, 20])
def test_a_box_with_only_a_border_is_drawn(store, angle):
    manifest = prepared(store)
    box = {"x0": 420.0, "y0": 690.0, "x1": 560.0, "y1": 740.0}
    _, cover = add_block(
        manifest, 1, bbox=box, text="", fill="#ffffff", store=store, target_lang="es",
    )
    _edit(manifest, cover, store, fill="none", border_color="#0000ff", border_width=2)
    if angle:
        _edit(manifest, cover, store, angle=angle)
    assert cover.manual_fill is None and cover.drawn_layout.border is not None
    assert (cover.drawn_layout.quad is not None) == bool(angle)
    if angle:
        assert _close(cover.manual_bbox, Rect(**box))
    final, plate, _ = _render(manifest, store)
    blue = _blue(final)
    corners = cover.drawn_layout.quad or frame_corners(Rect(**box), 0, (0, 0))
    for (ax, ay), (bx, by) in zip(corners, corners[1:] + corners[:1]):
        mx, my = (ax + bx) / 2, (ay + by) / 2
        assert blue[_px(my) - 1:_px(my) + 2, _px(mx) - 1:_px(mx) + 2].any()
    cx, cy = Rect(**box).center
    assert not blue[_px(cy) - 10:_px(cy) + 10, _px(cx) - 10:_px(cx) + 10].any()


@pytest.mark.parametrize("angle", [None, 30])
def test_the_sprite_holds_the_whole_border(store, angle):
    from ct_engine.api.routes import _block_out

    manifest, obj = _styled(store, border_color="#ff0000", border_width=6)
    if angle:
        _edit(manifest, obj, store, angle=angle)
    out = _block_out(obj)
    sprite = np.asarray(Image.open(io.BytesIO(render_sprite(obj, scale=2.0))).convert("RGBA"))
    area = Rect(**out["sprite_bbox"])
    assert sprite.shape[1] == pytest.approx(area.width * 2, abs=1)
    assert sprite.shape[0] == pytest.approx(area.height * 2, abs=1)
    # The border reaches 3pt past the frame; the picture is 8pt bigger.
    alpha = sprite[:, :, 3]
    assert not alpha[0].any() and not alpha[-1].any()
    assert not alpha[:, 0].any() and not alpha[:, -1].any()
    assert _red(sprite[:, :, :3])[alpha > 200].sum() > 500


# --- 2. Inside space ---------------------------------------------------------
def test_inset_keeps_room_for_the_words():
    inner, px, py = inset(Rect(x0=0, y0=0, x1=100, y1=10), 8.0)
    assert (px, py) == (8.0, 3.0)
    assert inner.height == pytest.approx(MIN_INNER)
    assert inset(Rect(x0=0, y0=0, x1=50, y1=50), None) == (Rect(x0=0, y0=0, x1=50, y1=50), 0.0, 0.0)
    assert inset(Rect(x0=0, y0=0, x1=3, y1=3), 2.0)[1:] == (0.0, 0.0)


@pytest.mark.parametrize("alignment", ["LEFT", "RIGHT", "JUSTIFY"])
def test_every_row_stays_inside_the_space(store, alignment):
    _, obj = _styled(store, padding=8, alignment=alignment)
    assert obj.manual_padding == 8.0
    rows = [r for r in obj.drawn_layout.rows if r.text]
    assert len(rows) > 2
    for r in rows:
        assert r.x0 >= BOX["x0"] + 8 - 0.01 and r.x1 <= BOX["x1"] - 8 + 0.01
    spread = [r for r in rows if r.spread]
    if alignment == "JUSTIFY":
        assert spread
        for r in spread:
            assert r.x0 == pytest.approx(BOX["x0"] + 8, abs=0.01)
            assert r.x1 == pytest.approx(BOX["x1"] - 8, abs=0.01)
    # The frame is still her whole box.
    assert _close(obj.drawn_bbox, Rect(**BOX))


def test_fit_leaves_the_space_empty(store):
    _, plain = _styled(store, size=None)
    _, spaced = _styled(store, size=None, padding=8)
    assert spaced.drawn_layout.size < plain.drawn_layout.size


def test_bottom_sits_on_the_inside_space(store):
    _, obj = _styled(store, padding=8, vertical_align="bottom")
    size = obj.drawn_layout.size
    last = [r.baseline for r in obj.drawn_layout.rows if r.text][-1]
    assert last == pytest.approx(BOX["y1"] - 8 - 0.2 * size, abs=0.01)
    _, top = _styled(store, padding=8)
    first = [r.baseline for r in top.drawn_layout.rows if r.text][0]
    assert first == pytest.approx(BOX["y0"] + 8 + 0.8 * top.drawn_layout.size, abs=0.01)


@pytest.mark.parametrize("where", ["top", "middle", "bottom"])
def test_a_spaced_box_handed_back_is_the_same_box(store, where):
    manifest, obj = _styled(store, padding=4, vertical_align=where)
    before = obj.drawn_bbox
    for _ in range(2):
        _edit(manifest, obj, store, bbox=obj.drawn_bbox.model_dump())
        assert _close(obj.drawn_bbox, before)


# --- 3. Turning --------------------------------------------------------------
def test_a_turned_printed_block_turns_about_its_centre(store):
    manifest, obj = _styled(store)
    frame = obj.drawn_bbox
    _edit(manifest, obj, store, angle=30)
    assert obj.manual_angle == 30.0 and _close(obj.manual_bbox, frame)
    layout = obj.drawn_layout
    assert layout.angle == 30.0
    want = frame_corners(frame, 30.0, frame.center)
    for (x, y), (x2, y2) in zip(want, layout.quad):
        assert x2 == pytest.approx(x, abs=0.01) and y2 == pytest.approx(y, abs=0.01)
    assert _close(obj.drawn_bbox, _aabb(layout.quad))
    _, _, out = _render(manifest, store)
    with fitz.open(stream=out, filetype="pdf") as doc:
        page = doc[0]
        text = page.get_text()
        dirs = [
            line["dir"]
            for block in page.get_text("dict")["blocks"]
            for line in block.get("lines", [])
            if "Instrucciones" in "".join(s["text"] for s in line["spans"])
        ]
    assert "Directions" not in text  # the original is still removed
    assert dirs and dirs[0][0] == pytest.approx(math.cos(math.radians(30)), abs=0.01)
    assert dirs[0][1] == pytest.approx(math.sin(math.radians(30)), abs=0.01)


def test_a_turned_block_is_the_same_ink_in_its_sprite(store):
    manifest, obj = _styled(store)
    _edit(manifest, obj, store, angle=-20, fill="#ffff00")
    final, plate, _ = _render(manifest, store)
    ink, (x0, y0) = sprite_on_page(obj)
    window = (slice(y0, y0 + ink.shape[0]), slice(x0, x0 + ink.shape[1]))
    changed = (np.abs(final - plate).max(axis=2) > 60)[window]
    ink = ink[:changed.shape[0], :changed.shape[1]]
    assert (ink & changed).sum() / max(1, (ink | changed).sum()) > 0.85


def test_turning_alone_pins_the_box_where_it_is(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    assert obj.manual_bbox is None
    frame = obj.drawn_bbox
    _edit(manifest, obj, store, angle=15)
    assert _close(obj.manual_bbox, frame)
    qx = sum(x for x, _ in obj.drawn_layout.quad) / 4
    qy = sum(y for _, y in obj.drawn_layout.quad) / 4
    assert (qx, qy) == pytest.approx(frame.center, abs=0.01)


def test_a_moved_and_turned_picture_label_still_covers_where_it_was(store):
    from test_tilted_text import _pixels

    manifest, label = _label_manifest(store, angle=1.0)
    assert label.oriented is None
    original = label.bbox
    moved = Rect(x0=original.x0, y0=original.y0 + 70, x1=original.x1, y1=original.y1 + 70)
    _edit(manifest, label, store, bbox=moved.model_dump(), angle=10)
    assert label.manual_angle == 10.0 and _close(label.manual_bbox, moved)
    src = store.get_bytes(manifest.source_key)
    from ct_engine.reconstruct.builder import rebuild_document

    out, _ = rebuild_document(src, manifest.pages, target_lang="es")
    with fitz.open(stream=out, filetype="pdf") as doc:
        final = _pixels(doc[0])
    spans = [s.bbox for line in label.lines for s in line.spans]
    for b in spans:
        patch = final[_px(b.y0) + 1:_px(b.y1) - 1, _px(b.x0) + 1:_px(b.x1) - 1]
        assert (patch.min(axis=2) > 215).mean() > 0.99
    assert original.y1 < _aabb(label.drawn_layout.quad).y0


def test_dragging_a_label_that_follows_its_tilt_keeps_the_slant(store):
    from ct_engine.api.routes import _block_out

    manifest, label = _label_manifest(store)
    out = _block_out(label)
    quad = label.drawn_layout.quad
    assert out["manual_angle"] is None and out["angle"] == pytest.approx(label.oriented.angle, abs=0.01)
    # The box the editor shows is the unturned frame.
    frame = turned_frame_of(label.drawn_layout)[0]
    assert _close(Rect(**out["render_bbox"]), frame)
    _edit(manifest, label, store, bbox=out["render_bbox"])
    assert label.manual_angle == pytest.approx(label.oriented.angle, abs=0.01)
    for (x, y), (x2, y2) in zip(quad, label.drawn_layout.quad):
        assert x2 == pytest.approx(x, abs=0.05) and y2 == pytest.approx(y, abs=0.05)
    again = _block_out(label)
    # The toolbar's "following" rule still holds.
    assert abs(again["angle"] - again["source_angle"]) < 0.05
    assert again["centred"] is True


def test_a_group_move_keeps_a_label_on_its_slant(store):
    from ct_engine.api.routes import _block_out

    manifest, label = _label_manifest(store)
    quad = label.drawn_layout.quad
    box = Rect(**_block_out(label)["render_bbox"])
    moved = Rect(x0=box.x0 + 10, y0=box.y0, x1=box.x1 + 10, y1=box.y1)
    apply_positions(
        manifest, 1, [{"id": label.id, "bbox": moved.model_dump()}],
        store=store, target_lang="es",
    )
    assert label.manual_angle == pytest.approx(label.oriented.angle, abs=0.01)
    for (x, y), (x2, y2) in zip(quad, label.drawn_layout.quad):
        assert x2 == pytest.approx(x + 10, abs=0.05) and y2 == pytest.approx(y, abs=0.05)


def test_straightening_a_tilted_label_is_straight(store):
    manifest, label = _label_manifest(store)
    _edit(manifest, label, store, angle=0)
    assert label.manual_rotation == 0.0 and label.manual_angle is None
    assert label.drawn_layout.quad is None


def test_straightening_a_box_she_turned_keeps_her_box(store):
    manifest, obj = _styled(store)
    _edit(manifest, obj, store, angle=30)
    frame = obj.manual_bbox
    _edit(manifest, obj, store, angle=0)
    assert obj.manual_angle is None and obj.manual_rotation is None
    assert _close(obj.manual_bbox, frame) and obj.drawn_layout.quad is None


def test_follow_the_tilt_clears_her_angle(store):
    manifest, label = _label_manifest(store)
    _edit(manifest, label, store, angle=40)
    _edit(manifest, label, store, rotation="follow")
    assert label.manual_angle is None and label.manual_bbox is None
    assert label.drawn_layout.angle == pytest.approx(label.oriented.angle, abs=0.01)
    _edit(manifest, label, store, angle=40)
    _edit(manifest, label, store, rotation="straight")
    assert label.manual_angle is None and label.manual_rotation == 0.0
    assert label.manual_bbox is not None  # her frame, now level


@pytest.mark.parametrize(
    "given, stored",
    [(370, 10.0), (-190, 170.0), (180, 180.0), (-180, 180.0), (540.004, 180.0),
     (0.004, None), (-359.999, None), (12.3456, 12.35), (-45, -45.0)],
)
def test_angles_are_stored_in_one_form(given, stored):
    assert _angle(given) == stored


@pytest.mark.parametrize("bad", [float("nan"), float("inf"), "left"])
def test_an_angle_that_cannot_be_read_is_refused(bad):
    with pytest.raises(EditError) as err:
        _angle(bad)
    assert err.value.message == "That angle could not be read."


def test_a_turned_frame_off_the_page_is_moved_not_shrunk():
    page = SimpleNamespace(bounds=Rect(x0=0, y0=0, x1=612, y1=792))
    box = Rect(x0=560, y0=-40, x1=700, y1=20)
    moved = _clamped(box, page, turned=True)
    assert (moved.width, moved.height) == (box.width, box.height)
    assert moved.center == pytest.approx((612.0, 0.0))
    assert _clamped(box, page, turned=False) == box.clamped_to(page.bounds)


def test_the_turned_frame_draws_the_same_quad_about_its_centre():
    """A widened picture label: its pivot is not its frame's centre."""
    rect = Rect(x0=100, y0=200, x1=260, y1=230)
    pivot = (150.0, 215.0)
    quad = frame_corners(rect, 13.0, pivot)
    layout = DrawnLayout(size=10, angle=13.0, pivot=pivot, quad=quad)
    frame, dx, dy = turned_frame_of(layout)
    assert (frame.width, frame.height) == pytest.approx((160, 30), abs=0.01)
    for (x, y), (x2, y2) in zip(quad, frame_corners(frame, 13.0, frame.center)):
        assert x2 == pytest.approx(x, abs=0.01) and y2 == pytest.approx(y, abs=0.01)
    # A point of the lines lands in the same place, shifted into the frame.
    for px_, py_ in ((110, 225), (250, 205)):
        a = _turn_point(px_, py_, pivot, 13.0)
        b = _turn_point(px_ + dx, py_ + dy, frame.center, 13.0)
        assert b == pytest.approx(a, abs=0.01)


def test_turned_sprites_get_a_new_address(store):
    from ct_engine.api.routes import _block_out

    _, obj = _styled(store)
    plain = sha256_bytes(
        (obj.drawn_layout.model_dump_json() + obj.drawn_bbox.model_dump_json()).encode()
    )[:12]
    assert _block_out(obj)["sprite_rev"] != plain


# --- 4. State -----------------------------------------------------------------
def test_an_old_manifest_loads_and_draws_the_same(store):
    manifest = prepared(store)
    before = {t.id: t.drawn_layout.lines for t in manifest.page(0).text_objects if t.drawn_layout}
    raw = manifest.model_dump(mode="json")
    for t in raw["pages"][0]["text_objects"]:
        for key in BOX_FIELDS:
            t.pop(key)
        if t.get("drawn_layout"):
            for key in ("border", "border_width", "border_dash", "centred"):
                t["drawn_layout"].pop(key)
    old = DocumentManifest.model_validate(raw)
    obj = old.page(0).text_objects[0]
    assert all(getattr(obj, f) is None for f in BOX_FIELDS)
    assert obj.drawn_layout.border is None and obj.drawn_layout.centred is False
    reconstruct_and_export(old, store=store, target_lang="es")
    after = {t.id: t.drawn_layout.lines for t in old.page(0).text_objects if t.drawn_layout}
    assert after == before


def test_each_box_style_is_stored_and_counts_as_hers(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    assert not obj.has_manual_layout
    for field, edit in (
        ("manual_padding", {"padding": 4}),
        ("manual_border_color", {"border_color": "#00FF00"}),
        ("manual_border_width", {"border_width": 0.5}),
        ("manual_border_dash", {"border_dash": "dotted"}),
        ("manual_angle", {"angle": 10}),
    ):
        _edit(manifest, obj, store, **edit)
        assert getattr(obj, field) is not None
        alone = TextObject.model_validate(
            obj.model_dump() | {"manual_bbox": None, "manual_border_color": None,
                                "manual_border_width": None, "manual_border_dash": None,
                                "manual_padding": None, "manual_angle": None,
                                field: getattr(obj, field)}
        )
        assert alone.has_manual_layout
    assert obj.manual_border_color == "#00ff00"


def test_a_weight_or_dash_with_no_border_gives_a_black_one(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    _edit(manifest, obj, store, border_width=3)
    assert obj.manual_border_color == "#000000" and obj.drawn_layout.border == (0, 0, 0)
    _edit(manifest, obj, store, border_color="none")
    assert obj.manual_border_color is None and obj.manual_border_width == 3.0
    assert obj.drawn_layout.border is None
    _edit(manifest, obj, store, border_dash="dashed", border_color="#ff0000")
    assert obj.manual_border_color == "#ff0000" and obj.manual_border_dash == "dashed"
    _edit(manifest, obj, store, border_dash="solid", padding=0)
    assert obj.manual_border_dash is None and obj.manual_padding is None


@pytest.mark.parametrize(
    "edit, message",
    [
        ({"border_color": "#ggg"}, "That border colour could not be read."),
        ({"border_width": 5}, "That border weight is not available."),
        ({"border_dash": "wavy"}, "That border style is not available."),
        ({"padding": 3}, "That inside space is not available."),
        ({"angle": float("nan")}, "That angle could not be read."),
    ],
)
def test_a_box_style_that_cannot_be_read_is_refused_and_records_nothing(store, edit, message):
    manifest = prepared(store)
    obj = _directions(manifest)
    history = len(manifest.page(0).edit_history)
    with pytest.raises(EditError) as err:
        _edit(manifest, obj, store, **edit)
    assert err.value.message == message
    with pytest.raises(EditError):
        apply_bulk(manifest, 1, [obj.id], store=store, target_lang="es", **edit)
    assert len(manifest.page(0).edit_history) == history


def test_undo_and_redo_step_through_each_box_style(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    steps = [
        ("manual_padding", {"padding": 8}, 8.0),
        ("manual_border_color", {"border_color": "#ff0000"}, "#ff0000"),
        ("manual_border_width", {"border_width": 6}, 6.0),
        ("manual_border_dash", {"border_dash": "dashed"}, "dashed"),
        ("manual_angle", {"angle": -25}, -25.0),
    ]
    for _, edit, _ in steps:
        _edit(manifest, obj, store, **edit)
    for field, _, value in reversed(steps):
        assert getattr(obj, field) == value
        undo(manifest, 1, store=store, target_lang="es")
        assert getattr(obj, field) is None
    assert obj.drawn_layout.quad is None and obj.drawn_layout.border is None
    for field, _, value in steps:
        redo(manifest, 1, store=store, target_lang="es")
        assert getattr(obj, field) == value
    assert obj.drawn_layout.angle == -25.0


def test_undo_to_a_snapshot_from_before_box_styles_clears_them(store):
    manifest = prepared(store)
    page = manifest.page(0)
    obj = _directions(manifest)
    old = _snapshot(page)
    for saved in old["blocks"].values():
        for key in BOX_FIELDS:
            saved.pop(key)
    page.edit_history.append(old)
    obj.manual_angle, obj.manual_padding = 30.0, 4.0
    obj.manual_border_color, obj.manual_border_width, obj.manual_border_dash = "#ff0000", 2.0, "dotted"
    undo(manifest, 1, store=store, target_lang="es")
    assert all(getattr(obj, f) is None for f in BOX_FIELDS)


@pytest.mark.parametrize("edited_text", [False, True])
def test_reset_clears_every_box_style(store, edited_text):
    manifest = prepared(store)
    obj = _directions(manifest)
    if edited_text:
        apply_edit(manifest, 1, obj.id, "Otro texto", store=store, target_lang="es")
    _edit(manifest, obj, store, padding=4, border_color="#ff0000", border_width=2,
          border_dash="dotted", angle=12)
    reset_edit(manifest, 1, obj.id, store=store, target_lang="es")
    assert all(getattr(obj, f) is None for f in BOX_FIELDS)
    assert obj.manual_bbox is None and obj.drawn_layout.quad is None
    assert obj.drawn_layout.border is None


def test_bulk_turns_each_box_about_its_own_centre(store):
    manifest = prepared(store)
    blocks = [t for t in manifest.page(0).text_objects if t.translated_text][:3]
    blocks[2].protected = True
    frames = [b.drawn_bbox for b in blocks]
    apply_bulk(
        manifest, 1, [b.id for b in blocks], border_color="#0000FF", border_width=2,
        padding=2, angle=10, store=store, target_lang="es",
    )
    for b, frame in zip(blocks[:2], frames):
        assert b.manual_border_color == "#0000ff" and b.manual_padding == 2.0
        assert b.manual_angle == 10.0 and _close(b.manual_bbox, frame)
        qx = sum(x for x, _ in b.drawn_layout.quad) / 4
        qy = sum(y for _, y in b.drawn_layout.quad) / 4
        assert (qx, qy) == pytest.approx(frame.center, abs=0.01)
    assert blocks[2].manual_angle is None and blocks[2].manual_border_color is None


def test_her_angle_wins_over_a_late_slant_and_says_nothing(store):
    manifest, label = _label_manifest(store)
    measured = label.oriented
    label.oriented = None
    label.orientation_checked = False
    label.manual_bbox = turned_frame_of(label.drawn_layout)[0]
    label.manual_angle = 25.0
    reconstruct_and_export(manifest, store=store, target_lang="es")
    assert label.oriented is not None and abs(label.oriented.angle - measured.angle) < 1
    assert label.manual_rotation is None and label.drawn_layout.angle == 25.0
    codes = {f.code for f in manifest.page(0).qa.findings}
    assert not codes & {"image_text_turned", "image_text_tilt_kept_straight"}


def test_the_editor_reads_every_box_style(store):
    from ct_engine.api.routes import _block_out

    manifest, obj = _styled(store)
    out = _block_out(obj)
    assert out["border_color"] is None and out["border_width"] == 1.0
    assert out["border_dash"] == "solid" and out["padding"] == 0
    assert out["manual_angle"] is None and out["centred"] is False
    _edit(manifest, obj, store, border_color="#ff0000", border_dash="dotted",
          padding=4, angle=30)
    out = _block_out(obj)
    assert (out["border_color"], out["border_width"], out["border_dash"]) == ("#ff0000", 1.0, "dotted")
    assert out["padding"] == 4.0 and out["manual_angle"] == 30.0 and out["angle"] == 30.0
    assert _close(Rect(**out["render_bbox"]), obj.manual_bbox)
    # The sprite is drawn unturned around the frame, with room for the border.
    assert _close(Rect(**out["sprite_bbox"]), obj.manual_bbox.expanded(3.0))


def test_a_turned_picture_labels_starting_size_is_its_letters(store):
    from ct_engine.api.routes import _block_out

    manifest, label = _label_manifest(store)
    label.style = label.style.model_copy(update={"font_size": 50.0})
    out = _block_out(label)
    assert out["auto_base_size"] == round(max(4.0, label.oriented.height * 0.92), 1)
    assert out["frame"]["height"] == pytest.approx(label.oriented.height, abs=0.01)


# --- 5. API -------------------------------------------------------------------
def test_the_editor_can_set_and_read_each_box_style(client):
    from test_api import _translate as translate, _upload, _wait_for_stage, _wait_terminal

    doc_id = _upload(client, "01_clean_worksheet.pdf")["document_id"]
    _wait_for_stage(client, doc_id)
    assert translate(client, doc_id).status_code in (200, 202)
    assert _wait_terminal(client, doc_id)["stage"] == "COMPLETE"
    blocks = client.get(f"/api/documents/{doc_id}/blocks").json()["pages"][0]["blocks"]
    block = next(b for b in blocks if b["original"].startswith("Directions"))
    other = next(b for b in blocks if b["id"] != block["id"])
    assert block["border_color"] is None and block["padding"] == 0
    assert block["manual_angle"] is None and block["angle"] == 0

    url = f"/api/documents/{doc_id}/pages/1/blocks/{block['id']}"
    r = client.patch(url, json={"border_color": "#FF0000", "border_width": 2,
                                "border_dash": "dashed", "padding": 4})
    assert r.status_code == 200, r.text
    styled = r.json()["block"]
    assert (styled["border_color"], styled["border_width"], styled["border_dash"]) == (
        "#ff0000", 2.0, "dashed")
    assert styled["padding"] == 4.0 and styled["sprite_rev"] != block["sprite_rev"]

    r = client.patch(url, json={"angle": 380})
    assert r.status_code == 200, r.text
    turned = r.json()["block"]
    assert turned["manual_angle"] == 20.0 and turned["angle"] == 20.0
    frame = turned["render_bbox"]
    assert frame == pytest.approx(styled["render_bbox"], abs=0.01)

    # Dragged with its angle: the box is the frame.
    moved = {k: v + 10 for k, v in frame.items()}
    r = client.patch(url, json={"bbox": moved, "angle": 20})
    assert r.status_code == 200, r.text
    assert r.json()["block"]["render_bbox"] == pytest.approx(moved, abs=0.01)

    for body, words in (({"angle": "NaN"}, "angle"), ({"border_width": 5}, "border weight"),
                        ({"padding": 3}, "inside space"), ({"border_dash": "wavy"}, "border style")):
        r = client.patch(url, json=body)
        assert r.status_code == 400 and words in r.text, r.text

    r = client.patch(
        f"/api/documents/{doc_id}/pages/1/blocks",
        json={"object_ids": [block["id"], other["id"]], "angle": -15, "padding": 2},
    )
    assert r.status_code == 200, r.text
    again = client.get(f"/api/documents/{doc_id}/blocks").json()["pages"][0]["blocks"]
    for b in again:
        if b["id"] in (block["id"], other["id"]):
            assert b["manual_angle"] == -15.0 and b["padding"] == 2.0

    r = client.post(f"/api/documents/{doc_id}/pages/1/undo")
    assert r.status_code == 200
    r = client.post(f"{url}/reset")
    assert r.status_code == 200
    reset = r.json()["block"]
    assert reset["border_color"] is None and reset["padding"] == 0
    assert reset["manual_angle"] is None and reset["angle"] == 0


# --- 6. Exports -----------------------------------------------------------------
def test_office_pictures_take_the_border_and_the_turned_frame(store):
    from ct_engine.pipeline.word_pictures import _labels_in

    manifest, label = _label_manifest(store, angle=1.0)
    page = manifest.page(0)
    image = page.image_objects[0]
    _edit(manifest, label, store, border_color="#ff0000", border_width=4)
    areas = _labels_in(page, image)
    assert any(isinstance(a, Rect) and _close(a, label.drawn_bbox.expanded(6.5)) for a in areas)
    _edit(manifest, label, store, angle=20)
    areas = _labels_in(page, image)
    quads = [a for a in areas if isinstance(a, list)]
    want = grow_quad(label.drawn_layout.quad, 6.5)
    assert any(
        all(abs(x - x2) < 0.01 and abs(y - y2) < 0.01 for (x, y), (x2, y2) in zip(q, want))
        for q in quads
    )
    # Where the original was covered is still copied.
    assert any(isinstance(a, Rect) and _close(a, label.bbox.expanded(2.5)) for a in areas)


def test_editable_powerpoint_keeps_a_turned_box_in_the_picture(store):
    from pptx import Presentation

    from ct_engine.reconstruct.layered_pptx import TURNED_TEXT_NOTE, export_layered_pptx

    manifest, obj = _styled(store, border_color="#ff0000", border_width=4)
    _edit(manifest, obj, store, angle=20)
    _, _, out = _render(manifest, store)
    deck = Presentation(io.BytesIO(export_layered_pptx(out)))
    slide = deck.slides[0]
    texts = " ".join(s.text for s in slide.shapes if s.has_text_frame)
    assert "Instrucciones" not in texts and texts.strip()
    assert TURNED_TEXT_NOTE in slide.notes_slide.notes_text_frame.text
    background = next(s for s in slide.shapes if s.shape_type == 13)
    pixels = np.asarray(Image.open(io.BytesIO(background.image.blob)).convert("RGB")).astype(int)
    assert _red(pixels).sum() > 200


def test_word_and_powerpoint_text_keep_their_own_box():
    """Deliberate: a block's border, space and turn are not written into
    the Office file's own text (see _apply_manual_style)."""
    from lxml import etree

    from ct_engine.reconstruct import docx_builder, pptx_builder

    plain = TextObject(
        id="t", bbox=Rect(x0=0, y0=0, x1=10, y1=10),
        norm_bbox=Rect(x0=0, y0=0, x1=10, y1=10).normalized(612, 792),
        manual_color="#ff0000", manual_font_size=12.0,
    )
    boxed = plain.model_copy(update={
        "manual_angle": 30.0, "manual_padding": 8.0, "manual_border_color": "#00ff00",
        "manual_border_width": 2.0, "manual_border_dash": "dashed",
        "manual_bbox": Rect(x0=0, y0=0, x1=10, y1=10),
    })
    w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for module, xml in ((docx_builder, f'<w:r xmlns:w="{w}"><w:t>x</w:t></w:r>'),
                        (pptx_builder, f'<a:r xmlns:a="{a}"><a:t>x</a:t></a:r>')):
        runs = [etree.fromstring(xml) for _ in range(2)]
        module._apply_manual_style(runs[0], plain)
        module._apply_manual_style(runs[1], boxed)
        assert etree.tostring(runs[0]) == etree.tostring(runs[1])


def test_the_original_is_removed_where_it_was_not_where_the_box_went(store):
    """The cover is about removing the source: turning and moving the
    translation never moves it."""
    manifest = prepared(store)
    obj = _directions(manifest)
    apply_edit(manifest, 1, obj.id, SPANISH, store=store, target_lang="es")
    _edit(manifest, obj, store, bbox=BOX, angle=45, padding=8)
    _, _, out = _render(manifest, store)
    with fitz.open(stream=out, filetype="pdf") as doc:
        words = doc[0].get_text("words", clip=fitz.Rect(*obj.bbox.as_tuple()))
    assert not [w for w in words if w[4].startswith("Directions")]
    assert obj.translation_status is TranslationStatus.MANUAL_EDIT
    assert obj.source is TextSource.NATIVE


@pytest.mark.parametrize("pad", [2.0, 4.0, 8.0])
def test_space_inside_an_automatic_box_grows_it_and_keeps_the_type(store, pad):
    # A box she has not placed grows outward by the space, as Slides grows a
    # shape round its text: a one-line label keeps its size and its width,
    # and nothing new overflows.
    manifest = prepared(store)
    for start in ("Name", "Living Environment"):
        obj = next(t for t in manifest.page(0).text_objects if t.flat_text.startswith(start))
        assert obj.manual_bbox is None
        size, before = obj.drawn_layout.size, obj.drawn_bbox
        _edit(manifest, obj, store, padding=pad)
        assert obj.manual_bbox is None and obj.drawn_layout.size == size
        after = obj.drawn_bbox
        assert after.x0 == pytest.approx(before.x0 - pad, abs=0.05)
        assert after.x1 == pytest.approx(before.x1 + pad, abs=0.05)
        assert after.y0 == pytest.approx(before.y0 - pad, abs=0.05)
        codes = {f.code for f in manifest.page(0).qa.findings if obj.id in (f.object_ids or [])}
        assert "text_overflow" not in codes


def test_the_border_weight_is_kept_while_it_is_off(store):
    from ct_engine.api.routes import _block_out

    manifest = prepared(store)
    obj = _directions(manifest)
    _edit(manifest, obj, store, border_color="#ff0000", border_width=6)
    _edit(manifest, obj, store, border_color="none")
    out = _block_out(obj)
    assert out["border_color"] is None and out["border_width"] == 6.0
