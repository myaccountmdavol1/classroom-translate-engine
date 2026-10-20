import sys, copy
from docx import Document
from docx.shared import Pt
from docx.oxml import parse_xml
out = sys.argv[1]
d = Document()
d.add_heading("Math and Science Review", 1)
def para(*parts):
    p = d.add_paragraph()
    for text, kind in parts:
        r = p.add_run(text)
        if kind == "sup": r.font.superscript = True
        if kind == "sub": r.font.subscript = True
        if kind == "sym":
            r.font.name = "Symbol"
            rpr = r._element.get_or_add_rPr()
            rpr.rFonts.set(  # Word's Insert Symbol stores w:sym; plain Symbol-font run is the common older form
                "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", "Symbol")
    return p
para(("1. A square has side x. Its area is x", ""), ("2", "sup"), (" square units. Explain why.", ""))
para(("2. Water is H", ""), ("2", "sub"), ("O and carbon dioxide is CO", ""), ("2", "sub"), (".", ""))
para(("3. The circumference of a circle is 2", ""), ("p", "sym"), ("r. Find it when r = 3 cm.", ""))
para(("4. Light travels about 3 × 10", ""), ("8", "sup"), (" meters per second.", ""))
para(("5. She finished in 1", ""), ("st", "sup"), (" place and he was 2", ""), ("nd", "sup"), (".", ""))
para(("6. Simplify the expression a", ""), ("m", "sup"), (" · a", ""), ("n", "sup"), (" and write your answer.", ""))
# 7: inline Office Math (OMML) fraction inside a sentence
p = d.add_paragraph(); p.add_run("7. Add ")
M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
p._element.append(parse_xml(f'<m:oMath {M}><m:f><m:num><m:r><m:t>3</m:t></m:r></m:num><m:den><m:r><m:t>4</m:t></m:r></m:den></m:f><m:r><m:t>+</m:t></m:r><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>8</m:t></m:r></m:den></m:f></m:oMath>'))
p.add_run(" and write the answer in simplest form.")
d.save(out)
