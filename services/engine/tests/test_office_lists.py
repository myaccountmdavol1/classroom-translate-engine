"""Lists and fonts in the Word and PowerPoint files.

A box that is one paragraph of the file becomes a real Word or PowerPoint
list; anywhere else each line starts with its marker, so no number is
lost. A catalog font is named in the file where it has the language's
letters. Each file is reopened with python-docx or python-pptx.
"""

from __future__ import annotations

import io
import re
import zipfile

import pytest
from lxml import etree

from conftest import corpus_file
from ct_engine.ingest.docx_parser import DocxParser
from ct_engine.model.document import DocumentManifest
from ct_engine.model.enums import TextSource, TranslationStatus
from ct_engine.pipeline.word_export import build_office_export
from ct_engine.ingest.pptx_walk import A, parse_xml
from ct_engine.reconstruct.pptx_builder import rebuild_pptx
from test_pptx import DECK, _pages, _slide_xml, _translate

TITLE = "Cells and Energy: Review Packet"


def _strip_numbering(data: bytes) -> bytes:
    """The worksheet without its numbering part, its relationship or its
    content type: a document that never had a list."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename == "word/numbering.xml":
                continue
            body = src.read(item.filename)
            if item.filename in ("word/_rels/document.xml.rels", "[Content_Types].xml"):
                body = re.sub(rb"<(Relationship|Override)[^>]*numbering[^>]*/>", b"", body)
            dst.writestr(item, body)
    return out.getvalue()


def _number_the_title(data: bytes) -> bytes:
    """The worksheet with Word itself numbering the title paragraph."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "word/document.xml":
                xml = body.decode()
                at = xml.index(TITLE)
                p = xml.rindex("<w:p>", 0, at) if "<w:p>" in xml[:at] else xml.rindex("<w:p ", 0, at)
                end = xml.index(">", p) + 1
                ppr = '<w:pPr><w:numPr><w:ilvl w:val="0"/><w:numId w:val="1"/></w:numPr></w:pPr>'
                if xml[end:end + 7] == "<w:pPr>":
                    xml = xml[:end + 7] + ppr[7:-8] + xml[end + 7:]
                else:
                    xml = xml[:end] + ppr + xml[end:]
                body = xml.encode()
            dst.writestr(item, body)
    return out.getvalue()


def _export(store, *, data=None, lang="es", **styles):
    data = data or corpus_file("20_worksheet.docx")
    store.put_bytes("documents/w/source.docx", data)
    manifest = DocumentManifest(
        document_id="w", filename="w.docx", content_type="x", byte_size=len(data),
        sha256="0", source_key="documents/w/source.docx",
    )
    manifest.office_source_key, manifest.office_format = "documents/w/source.docx", "docx"
    with DocxParser().open_bytes(data) as doc:
        manifest.pages = [doc.parse_page(i) for i in range(doc.page_count)]
    for p in manifest.pages:
        for i, t in enumerate(p.text_objects):
            t.source, t.reading_index = TextSource.NATIVE, i
            t.translated_text = f"<{t.flat_text}>"
            t.translation_status = TranslationStatus.TRANSLATED
    title = next(t for t in manifest.pages[0].text_objects if t.text.strip() == TITLE)
    text = styles.pop("text", None)
    if text is not None:
        title.translated_text = text
        title.manually_edited = True
    for name, value in styles.items():
        setattr(title, name, value)
    return build_office_export(manifest, store=store, target_lang=lang).data


def _xml(data: bytes, part: str) -> str:
    return zipfile.ZipFile(io.BytesIO(data)).read(part).decode()


def _title_paragraph(data: bytes, needle: str) -> str:
    xml = _xml(data, "word/document.xml")
    at = xml.index(needle)
    return xml[max(xml.rindex("<w:p>", 0, at) if "<w:p>" in xml[:at] else 0,
                   xml.rindex("<w:p ", 0, at) if "<w:p " in xml[:at] else 0): xml.index("</w:p>", at)]


def test_a_one_paragraph_box_becomes_a_real_word_list(store):
    import docx

    data = _export(store, manual_list="number", manual_list_start=3)
    para = _title_paragraph(data, "&lt;Cells and Energy")
    assert "<w:numPr>" in para and '<w:ilvl w:val="0"/>' in para
    num_id = re.search(r'<w:numId w:val="(\d+)"/>', para).group(1)
    numbering = _xml(data, "word/numbering.xml")
    num = re.search(rf'<w:num w:numId="{num_id}">(.*?)</w:num>', numbering).group(1)
    assert '<w:startOverride w:val="3"/>' in num
    abstract_id = re.search(r'<w:abstractNumId w:val="(\d+)"/>', num).group(1)
    abstract = re.search(
        rf'<w:abstractNum w:abstractNumId="{abstract_id}">(.*?)</w:abstractNum>', numbering
    ).group(1)
    assert '<w:numFmt w:val="decimal"/>' in abstract and '<w:lvlText w:val="%1."/>' in abstract
    assert '<w:numFmt w:val="lowerLetter"/>' in abstract  # level 1, as the pages count
    # Every abstractNum still comes before the first num.
    assert numbering.rindex("<w:abstractNum ") < numbering.index("<w:num ")
    # The words carry no typed marker: Word draws the number.
    assert "&lt;Cells and Energy: Review Packet&gt;" in para
    docx.Document(io.BytesIO(data))


def test_a_document_with_no_numbering_part_gets_one(store):
    import docx

    data = _export(store, data=_strip_numbering(corpus_file("20_worksheet.docx")),
                   manual_list="bullet", manual_list_style="check")
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert "word/numbering.xml" in names
    assert 'relationships/numbering" Target="numbering.xml"' in _xml(data, "word/_rels/document.xml.rels")
    assert 'PartName="/word/numbering.xml"' in _xml(data, "[Content_Types].xml")
    numbering = _xml(data, "word/numbering.xml")
    assert '<w:numFmt w:val="bullet"/>' in numbering and 'w:val="✓"' in numbering
    assert 'w:ascii="Noto Sans Symbols 2"' in numbering
    document = docx.Document(io.BytesIO(data))
    assert document.part.numbering_part is not None


def test_a_paragraph_word_numbers_keeps_its_own_numbers(store):
    data = _export(store, data=_number_the_title(corpus_file("20_worksheet.docx")),
                   manual_list="number")
    para = _title_paragraph(data, "&lt;Cells and Energy")
    assert para.count("<w:numPr>") == 1 and '<w:numId w:val="1"/>' in para
    assert "1. &lt;Cells" not in para


def test_a_box_of_several_lines_starts_each_with_its_marker(store):
    data = _export(store, text="uno\ndos", manual_list="number")
    para = _title_paragraph(data, ">1. uno<")
    assert "<w:numPr>" not in para
    assert re.search(r">1\. uno</w:t>\s*<w:br/>\s*<w:t>2\. dos<", para)
    data = _export(store, text="uno\ndos", manual_list="bullet")
    assert "• uno" in _xml(data, "word/document.xml")


def test_word_names_a_font_that_has_the_language(store):
    data = _export(store, manual_font_family="andika")
    para = _title_paragraph(data, "&lt;Cells and Energy")
    assert 'w:ascii="Andika"' in para and 'w:hAnsi="Andika"' in para
    data = _export(store, lang="zh-Hans", manual_font_family="andika")
    assert "Andika" not in _xml(data, "word/document.xml")
    data = _export(store, manual_font_family="serif")  # the older styles stay unnamed
    assert 'w:ascii="Noto' not in _title_paragraph(data, "&lt;Cells and Energy")


def test_word_rfonts_stay_in_schema_order():
    from types import SimpleNamespace

    from ct_engine.ingest.docx_walk import W
    from ct_engine.reconstruct.docx_builder import _apply_manual_style

    run = etree.fromstring(
        f'<w:r xmlns:w="{W[1:-1]}"><w:rPr><w:b/><w:sz w:val="20"/></w:rPr><w:t>x</w:t></w:r>'
    )
    obj = SimpleNamespace(
        manual_color=None, manual_font_size=None, manual_underline=False,
        manual_highlight=None, manual_font_family="lexend",
    )
    _apply_manual_style(run, obj)
    names = [etree.QName(c).localname for c in run.find(f"{W}rPr")]
    assert names == ["rFonts", "b", "sz"]


def _slide_pages(**styles):
    pages = _translate(_pages(), lambda t: f"<{t}>")
    target = next(o for o in pages[0].text_objects if o.text.startswith("Plants use"))
    text = styles.pop("text", None)
    if text is not None:
        target.translated_text = text
    for name, value in styles.items():
        setattr(target, name, value)
    return pages


def _slide_paragraph(xml: str, needle: str) -> str:
    at = xml.index(needle)
    return xml[xml.rindex("<a:p>", 0, at): xml.index("</a:p>", at)]


def test_a_one_paragraph_box_becomes_real_powerpoint_numbers():
    from pptx import Presentation

    pages = _slide_pages(manual_list="number", manual_list_start=3)
    data = rebuild_pptx(corpus_file(DECK), pages).data
    para = _slide_paragraph(_slide_xml(data, 1), "&lt;Plants use")
    assert '<a:buAutoNum type="arabicPeriod" startAt="3"/>' in para
    assert 'lvl="0"' in para and 'indent="-' in para
    Presentation(io.BytesIO(data))


def test_powerpoint_bullets_and_levels():
    from pptx import Presentation

    pages = _slide_pages(manual_list="bullet", manual_indents=[1])
    data = rebuild_pptx(corpus_file(DECK), pages).data
    para = _slide_paragraph(_slide_xml(data, 1), "&lt;Plants use")
    assert '<a:buFont typeface="Noto Sans Symbols 2"/>' in para
    assert '<a:buChar char="◦"/>' in para and 'lvl="1"' in para
    marl = int(re.search(r'marL="(\d+)"', para).group(1))
    indent = int(re.search(r'indent="(-\d+)"', para).group(1))
    assert marl + indent > 0  # the marker starts a level in
    Presentation(io.BytesIO(data))


def test_powerpoint_lines_start_with_their_markers_where_a_list_cannot_go():
    pages = _slide_pages(text="uno\ndos", manual_list="number", manual_list_style="lower-alpha")
    xml = _slide_xml(rebuild_pptx(corpus_file(DECK), pages).data, 1)
    assert "a. uno" in xml and "b. dos" in xml
    para = _slide_paragraph(xml, "a. uno")
    assert "buAutoNum" not in para
    # The body placeholder inherits the master's "•" bullets; without an
    # explicit no-bullet PowerPoint drew "• a. uno".
    assert "<a:buNone/>" in para


def test_powerpoint_inherited_bullets_are_read_from_the_master():
    from ct_engine.reconstruct.pptx_builder import _Inherited

    with zipfile.ZipFile(io.BytesIO(corpus_file(DECK))) as zin:
        root = parse_xml(zin.read("ppt/slides/slide1.xml"))
        inherited = _Inherited(zin)
        paras = list(root.iter(f"{A}p"))
        body = next(p for p in paras if "Plants use" in "".join(p.itertext()))
        title = next(p for p in paras if "Photosynthesis" in "".join(p.itertext()))
        assert inherited.bullets("ppt/slides/slide1.xml", body) is True
        assert inherited.bullets("ppt/slides/slide1.xml", title) is False


def test_powerpoint_title_lines_get_no_extra_bullet_setting():
    pages = _translate(_pages(), lambda t: f"<{t}>")
    title = next(o for o in pages[0].text_objects if o.text.startswith("Photosynthesis"))
    title.translated_text, title.manual_list = "uno\ndos", "bullet"
    xml = _slide_xml(rebuild_pptx(corpus_file(DECK), pages).data, 1)
    para = _slide_paragraph(xml, "• uno")
    assert "buNone" not in para


def test_powerpoint_names_the_font():
    from pptx import Presentation

    pages = _slide_pages(manual_font_family="andika")
    data = rebuild_pptx(corpus_file(DECK), pages).data
    xml = _slide_xml(data, 1)
    para = _slide_paragraph(xml, "&lt;Plants use")
    assert '<a:latin typeface="Andika"/>' in para
    rpr = para[para.index("<a:rPr"): para.index("</a:rPr>")]
    assert rpr.index("<a:latin") < rpr.index("<a:ea") < rpr.index("<a:cs")
    Presentation(io.BytesIO(data))


@pytest.mark.parametrize("style,expected", [
    ("decimal-paren", "arabicParenR"), ("upper-alpha", "alphaUcPeriod"),
    ("lower-roman", "romanLcPeriod"),
])
def test_powerpoint_number_styles(style, expected):
    pages = _slide_pages(manual_list="number", manual_list_style=style)
    para = _slide_paragraph(_slide_xml(rebuild_pptx(corpus_file(DECK), pages).data, 1), "&lt;Plants use")
    assert f'type="{expected}"' in para and "startAt" not in para
