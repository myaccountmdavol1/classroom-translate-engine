"""Google Slides' text styles: where the words sit up and down in the box,
justified text, line spacing, underline and highlight.

Each is a choice she makes in the editor's toolbar. The engine draws it in
the PDF and in the block's live picture alike, keeps it through undo and
Reset, and carries what Word and PowerPoint can hold into those files.
"""

from __future__ import annotations

import io
import zipfile
from types import SimpleNamespace

import fitz
import numpy as np
import pytest
from PIL import Image, ImageFilter

from conftest import corpus_file
from ct_engine.model.enums import TextAlignment, TranslationStatus
from ct_engine.model.objects import DrawnLayout, TextObject, grow_quad
from ct_engine.pipeline.edit import (
    EditError,
    _snapshot,
    apply_bulk,
    apply_edit,
    apply_layout,
    redo,
    reset_edit,
    undo,
)
from ct_engine.reconstruct.builder import rebuild_document
from ct_engine.reconstruct.fitting import (
    LINE_SPACING_PRESETS,
    fit_text,
    fit_to_box,
    justifiable,
    wrap_with_ends,
)
from ct_engine.reconstruct.sprite import PAD, render_sprite
from ct_engine.fonts.registry import REGISTRY
from ct_engine.model.geometry import Rect
from test_edit import prepared
from test_reconstruct import prepared_pages
from test_translate_pipeline import EchoProvider

SPANISH = (
    "Instrucciones: lea cada pregunta con cuidado y escriba sus respuestas "
    "en oraciones completas usando el vocabulario de la unidad.\n"
    "Use evidencia del texto para apoyar cada respuesta que escriba."
)
# A box with room under the words, clear of everything else drawn.
BOX = {"x0": 72.0, "y0": 420.0, "x1": 330.0, "y1": 560.0}
SIZE = 11.0
STYLE_FIELDS = (
    "manual_vertical_align", "manual_line_spacing", "manual_underline", "manual_highlight",
)


def _directions(manifest):
    return next(
        t for t in manifest.page(0).text_objects if t.flat_text.startswith("Directions")
    )


def _styled(store, *, text=SPANISH, box=BOX, size=SIZE, lang="es", **edits):
    """The Directions block with `text`, placed in `box` at a locked size,
    then given `edits`."""
    manifest = prepared(store)
    obj = _directions(manifest)
    apply_edit(manifest, 1, obj.id, text, store=store, target_lang=lang)
    layout = {"font_size": size} if size else {"fit": True}
    apply_layout(manifest, 1, obj.id, bbox=box, store=store, target_lang=lang, **layout)
    if edits:
        apply_layout(manifest, 1, obj.id, store=store, target_lang=lang, **edits)
    return manifest, obj


def _baselines(obj):
    return [r.baseline for r in obj.drawn_layout.rows if r.text]


def _font(layout):
    if layout.font_builtin:
        return fitz.Font(layout.font_builtin)
    return fitz.Font(fontfile=layout.font_path)


def _page_pixels(manifest, store, dpi=96, plates=None):
    src = store.get_bytes(manifest.source_key)
    on_plate = None
    if plates is not None:
        def on_plate(index, page):
            plates[index] = _rgb(page.get_pixmap(dpi=dpi, alpha=False))
    out, _ = rebuild_document(src, manifest.pages, target_lang="es", on_plate=on_plate)
    with fitz.open(stream=out, filetype="pdf") as doc:
        return _rgb(doc[0].get_pixmap(dpi=dpi, alpha=False)), out


def _rgb(pix):
    return np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3).astype(int)


def _px(v, dpi=96):
    return int(round(v * dpi / 72))


# --- 1. Compatibility -----------------------------------------------------
def test_an_old_manifest_and_layout_load_with_nothing_styled(store):
    manifest = prepared(store)
    raw = _directions(manifest).model_dump(mode="json")
    for key in STYLE_FIELDS:
        raw.pop(key)
    for key in ("leading", "rows", "underline", "highlight"):
        raw["drawn_layout"].pop(key)
    old = TextObject.model_validate(raw)
    assert old.manual_vertical_align is None and old.manual_line_spacing is None
    assert old.manual_underline is False and old.manual_highlight is None
    assert not old.has_manual_layout
    layout = DrawnLayout.model_validate({"size": 10.0, "lines": [[1.0, 2.0, "hola"]]})
    assert layout.rows == [] and layout.leading is None
    assert layout.underline is False and layout.highlight is None


def test_an_old_layout_draws_its_sprite_without_decorations():
    layout = DrawnLayout(font_builtin="helv", size=12.0, lines=[(10.0, 20.0, "Hola")])
    obj = TextObject(
        id="t", bbox=Rect(x0=0, y0=0, x1=80, y1=30),
        norm_bbox=Rect(x0=0, y0=0, x1=80, y1=30).normalized(612, 792),
        drawn_layout=layout, drawn_bbox=Rect(x0=0, y0=0, x1=80, y1=30),
    )
    assert render_sprite(obj)[:4] == b"\x89PNG"


def test_undo_to_a_snapshot_from_before_these_styles_clears_them(store):
    manifest = prepared(store)
    page = manifest.page(0)
    obj = _directions(manifest)
    old = _snapshot(page)
    for saved in old["blocks"].values():
        for key in STYLE_FIELDS:
            saved.pop(key)
    page.edit_history.append(old)
    obj.manual_underline = True
    obj.manual_highlight = "#fff59d"
    obj.manual_vertical_align = "bottom"
    obj.manual_line_spacing = 2.0
    undo(manifest, 1, store=store, target_lang="es")
    assert not obj.manual_underline and obj.manual_highlight is None
    assert obj.manual_vertical_align is None and obj.manual_line_spacing is None
    # Its other fields are taken from the snapshot as before.
    assert obj.translation_status is TranslationStatus.TRANSLATED


# --- 2. Edits -------------------------------------------------------------
def test_each_style_is_stored_in_its_plain_form(store):
    manifest = prepared(store)
    obj = _directions(manifest)

    def edit(**kw):
        apply_layout(manifest, 1, obj.id, store=store, target_lang="es", **kw)

    edit(vertical_align="middle", line_spacing=1.1500001, underline=True, highlight="#FFF59D")
    assert obj.manual_vertical_align == "middle"
    assert obj.manual_line_spacing == 1.15
    assert obj.manual_underline is True
    assert obj.manual_highlight == "#fff59d"
    assert obj.has_manual_layout
    assert obj.manually_edited and obj.translation_status is TranslationStatus.MANUAL_EDIT
    edit(vertical_align="top", line_spacing=0, underline=False, highlight="none")
    assert obj.manual_vertical_align is None and obj.manual_line_spacing is None
    assert obj.manual_underline is False and obj.manual_highlight is None


@pytest.mark.parametrize(
    "edit, message",
    [
        ({"vertical_align": "side"}, "That vertical alignment could not be read."),
        ({"line_spacing": 1.3}, "That line spacing is not available."),
        ({"highlight": "#ggg"}, "That highlight colour could not be read."),
    ],
)
def test_a_style_that_cannot_be_read_is_refused_and_records_nothing(store, edit, message):
    manifest = prepared(store)
    obj = _directions(manifest)
    history = len(manifest.page(0).edit_history)
    with pytest.raises(EditError) as err:
        apply_layout(manifest, 1, obj.id, store=store, target_lang="es", **edit)
    assert err.value.message == message
    assert len(manifest.page(0).edit_history) == history
    with pytest.raises(EditError):
        apply_bulk(manifest, 1, [obj.id], store=store, target_lang="es", **edit)
    assert len(manifest.page(0).edit_history) == history


def test_bulk_styles_every_selected_block_but_a_locked_one(store):
    manifest = prepared(store)
    blocks = [t for t in manifest.page(0).text_objects if t.translated_text][:3]
    blocks[2].protected = True
    apply_bulk(
        manifest, 1, [b.id for b in blocks], vertical_align="bottom", line_spacing=1.5,
        underline=True, highlight="#B9F6CA", store=store, target_lang="es",
    )
    for b in blocks[:2]:
        assert (b.manual_vertical_align, b.manual_line_spacing) == ("bottom", 1.5)
        assert b.manual_underline and b.manual_highlight == "#b9f6ca"
        assert b.drawn_layout.underline
    assert not blocks[2].manual_underline and blocks[2].manual_highlight is None


@pytest.mark.parametrize("edited_text", [False, True])
def test_reset_clears_every_style(store, edited_text):
    manifest = prepared(store)
    obj = _directions(manifest)
    if edited_text:
        apply_edit(manifest, 1, obj.id, "Otro texto", store=store, target_lang="es")
    apply_layout(
        manifest, 1, obj.id, vertical_align="middle", line_spacing=2.0, underline=True,
        highlight="#fff59d", store=store, target_lang="es",
    )
    reset_edit(manifest, 1, obj.id, store=store, target_lang="es")
    assert obj.manual_vertical_align is None and obj.manual_line_spacing is None
    assert obj.manual_underline is False and obj.manual_highlight is None
    assert not obj.drawn_layout.underline and obj.drawn_layout.highlight is None


def test_undo_and_redo_step_through_each_style(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    steps = [
        ("manual_vertical_align", {"vertical_align": "bottom"}, "bottom"),
        ("manual_line_spacing", {"line_spacing": 2.0}, 2.0),
        ("manual_underline", {"underline": True}, True),
        ("manual_highlight", {"highlight": "#e1bee7"}, "#e1bee7"),
    ]
    for _, edit, _ in steps:
        apply_layout(manifest, 1, obj.id, store=store, target_lang="es", **edit)
    for field, _, value in reversed(steps):
        assert getattr(obj, field) == value
        undo(manifest, 1, store=store, target_lang="es")
        assert getattr(obj, field) == TextObject.model_fields[field].get_default()
    for field, _, value in steps:
        redo(manifest, 1, store=store, target_lang="es")
        assert getattr(obj, field) == value


# --- 3. API ---------------------------------------------------------------
def test_the_editor_can_set_and_read_each_style(client):
    from test_api import _translate, _upload, _wait_for_stage, _wait_terminal

    doc_id = _upload(client, "01_clean_worksheet.pdf")["document_id"]
    _wait_for_stage(client, doc_id)
    assert _translate(client, doc_id).status_code in (200, 202)
    assert _wait_terminal(client, doc_id)["stage"] == "COMPLETE"
    blocks = client.get(f"/api/documents/{doc_id}/blocks").json()["pages"][0]["blocks"]
    block = next(b for b in blocks if b["original"].startswith("Directions"))
    assert block["vertical_align"] == "top" and block["line_spacing"] is None
    assert block["underline"] is False and block["highlight"] is None
    assert block["justifiable"] is True and block["auto_line_height"] > 1.0
    assert block["rows"] and all({"text", "end", "spread"} <= set(r) for r in block["rows"])
    assert block["rows"][-1]["end"] is True

    url = f"/api/documents/{doc_id}/pages/1/blocks/{block['id']}"
    r = client.patch(url, json={"underline": True})
    assert r.status_code == 200, r.text
    underlined = r.json()["block"]
    assert underlined["underline"] is True
    assert underlined["sprite_rev"] != block["sprite_rev"]

    r = client.patch(url, json={"vertical_align": "bottom", "line_spacing": 1.5})
    assert r.status_code == 200, r.text
    moved = r.json()["block"]
    assert moved["vertical_align"] == "bottom" and moved["manual_vertical_align"] == "bottom"
    assert moved["line_spacing"] == 1.5 and moved["line_height"] == pytest.approx(1.5, abs=0.01)
    assert moved["sprite_rev"] != underlined["sprite_rev"]

    r = client.patch(url, json={"line_spacing": 1.3})
    assert r.status_code == 400 and "line spacing" in r.text

    r = client.patch(
        f"/api/documents/{doc_id}/pages/1/blocks",
        json={"object_ids": [block["id"]], "highlight": "#fff59d"},
    )
    assert r.status_code == 200, r.text
    again = client.get(f"/api/documents/{doc_id}/blocks").json()["pages"][0]["blocks"]
    assert next(b for b in again if b["id"] == block["id"])["highlight"] == "#fff59d"

    r = client.post(f"{url}/reset")
    assert r.status_code == 200
    reset = r.json()["block"]
    assert reset["underline"] is False and reset["highlight"] is None
    assert reset["vertical_align"] == "top" and reset["line_spacing"] is None


def test_a_justified_source_reads_as_left_because_it_is_drawn_left(store):
    from ct_engine.api.routes import _block_out

    manifest = prepared(store)
    obj = _directions(manifest)
    obj.alignment = TextAlignment.JUSTIFY
    assert _block_out(obj)["alignment"] == "LEFT"
    obj.manual_alignment = TextAlignment.JUSTIFY
    assert _block_out(obj)["alignment"] == "JUSTIFY"


def test_line_height_is_what_the_picture_used(store):
    from ct_engine.api.routes import _block_out

    _, obj = _styled(store, line_spacing=2.0)
    out = _block_out(obj)
    assert out["line_height"] == pytest.approx(2.0, abs=0.01)
    assert out["auto_line_height"] == pytest.approx(1.18)  # size is locked


def test_automatic_spacing_is_measured_against_the_size_sizing_starts_from(store):
    from ct_engine.api.routes import _line_heights

    manifest = prepared(store)
    obj = _directions(manifest)
    obj.style.font_size = 10.0
    obj.line_spacing = 13.0
    assert _line_heights(obj)[1] == pytest.approx(1.3)
    # A size she set without locking it is where lay_out starts (base_size).
    obj.manual_font_size = 11.0
    assert _line_heights(obj)[1] == pytest.approx(13.0 / 11.0, abs=1e-4)


def test_a_layout_saved_before_its_leading_reads_it_off_the_baselines(store):
    from ct_engine.api.routes import _block_out

    _, obj = _styled(store, line_spacing=1.5)
    old = obj.drawn_layout.model_dump()
    for key in ("leading", "rows"):
        old.pop(key)
    obj.drawn_layout = DrawnLayout.model_validate(old)
    assert _block_out(obj)["line_height"] == pytest.approx(1.5, abs=0.01)


def test_automatic_blocks_say_where_their_sizing_starts(store):
    from ct_engine.api.routes import _block_out

    _, locked = _styled(store)
    assert _block_out(locked)["auto_base_size"] is None
    manifest = prepared(store)
    obj = _directions(manifest)
    assert _block_out(obj)["auto_base_size"] == obj.style.font_size
    obj.manual_font_size = 14.0
    assert _block_out(obj)["auto_base_size"] == 14.0
    obj.manual_fit = True
    assert _block_out(obj)["auto_base_size"] is None


# --- 4. Up and down -------------------------------------------------------
def test_top_keeps_the_first_baseline_where_it_always_was(store):
    _, obj = _styled(store)
    assert _baselines(obj)[0] == pytest.approx(BOX["y0"] + 0.8 * SIZE, abs=0.01)


def test_bottom_puts_the_last_baseline_a_descent_above_the_box_bottom(store):
    _, obj = _styled(store, vertical_align="bottom")
    assert _baselines(obj)[-1] == pytest.approx(BOX["y1"] - 0.2 * SIZE, abs=0.01)
    assert obj.drawn_bbox.y1 == pytest.approx(BOX["y1"], abs=0.01)


def test_middle_leaves_the_same_room_above_and_below(store):
    _, obj = _styled(store, vertical_align="middle")
    rows = _baselines(obj)
    above = rows[0] - 0.8 * SIZE - BOX["y0"]
    below = BOX["y1"] - rows[-1] - 0.2 * SIZE
    assert above == pytest.approx(below, abs=0.01)
    assert above > 10


@pytest.mark.parametrize("where, k", [("top", 0.0), ("middle", 0.5), ("bottom", 1.0)])
def test_text_too_long_for_its_box_stays_at_the_top(store, where, k):
    box = {**BOX, "y1": BOX["y0"] + 20}
    _, obj = _styled(store, box=box, vertical_align=where)
    step = obj.drawn_layout.leading
    assert _baselines(obj)[0] == pytest.approx(
        box["y0"] + 0.8 * SIZE + k * (step - SIZE), abs=0.01
    )


def test_fit_and_bottom_sit_on_the_box_bottom(store):
    _, obj = _styled(store, size=None, vertical_align="bottom")
    size = obj.drawn_layout.size
    assert _baselines(obj)[-1] == pytest.approx(BOX["y1"] - 0.2 * size, abs=0.01)


def test_a_box_grown_downward_aligns_within_its_new_height():
    data, pages = prepared_pages("01_clean_worksheet.pdf")
    obj = next(t for t in pages[0].text_objects if t.flat_text.startswith("Directions"))
    obj.translated_text = ("palabra " * 200)[:700]
    obj.manual_vertical_align = "bottom"
    rebuild_document(data, pages, target_lang="es")
    assert obj.drawn_bbox.y1 > obj.bbox.y1 + 1  # it grew
    assert _baselines(obj)[-1] == pytest.approx(
        obj.drawn_bbox.y1 - 0.2 * obj.drawn_layout.size, abs=0.01
    )


@pytest.mark.parametrize("where", ["middle", "bottom"])
def test_dragging_an_aligned_box_does_not_grow_it(store, where):
    manifest, obj = _styled(store, vertical_align=where)
    heights = []
    for _ in range(3):
        b = obj.drawn_bbox
        apply_layout(
            manifest, 1, obj.id, bbox={"x0": b.x0 + 5, "y0": b.y0 + 5, "x1": b.x1 + 5, "y1": b.y1 + 5},
            store=store, target_lang="es",
        )
        heights.append(obj.drawn_bbox.height)
    assert max(heights) - min(heights) <= 0.01
    assert heights[0] == pytest.approx(BOX["y1"] - BOX["y0"], abs=0.01)


# --- 5. Line spacing ------------------------------------------------------
@pytest.mark.parametrize("spacing", LINE_SPACING_PRESETS)
def test_each_line_spacing_is_that_many_times_the_size(store, spacing):
    _, obj = _styled(store, line_spacing=spacing)
    rows = [r.baseline for r in obj.drawn_layout.rows]
    gaps = np.diff(rows)
    assert len(gaps) >= 3
    assert np.allclose(gaps, spacing * SIZE, atol=0.02)


def test_automatic_text_keeps_her_spacing_and_shrinks_instead(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    apply_layout(manifest, 1, obj.id, line_spacing=2.0, store=store, target_lang="es")
    layout = obj.drawn_layout
    assert layout.leading == pytest.approx(2.0 * layout.size, abs=0.01)


def test_fit_to_box_counts_her_spacing():
    font = REGISTRY.resolve(SPANISH)
    rect = Rect(x0=0, y0=0, x1=200, y1=120)
    fit = fit_to_box(SPANISH, rect, font, leading=2.0)
    assert len(fit.lines) * 2.0 * fit.font_size <= rect.height + 1e-6
    assert fit.leading == pytest.approx(2.0 * fit.font_size, abs=0.02)
    plain = fit_to_box(SPANISH, rect, font)
    assert plain.font_size > fit.font_size


def test_fit_text_never_tightens_her_spacing():
    font = REGISTRY.resolve(SPANISH)
    fit = fit_text(SPANISH, Rect(x0=0, y0=0, x1=200, y1=60), font, base_size=11, leading=1.5)
    assert fit.leading >= 1.5 * fit.font_size - 0.01


# --- 6. Justify -----------------------------------------------------------
def test_justified_rows_run_from_edge_to_edge(store):
    _, obj = _styled(store, alignment="JUSTIFY")
    layout = obj.drawn_layout
    font = _font(layout)
    spread = [r for r in layout.rows if not r.end]
    assert len(spread) >= 2
    for row in spread:
        words = [e for e in layout.lines if abs(e[1] - row.baseline) < 0.01]
        assert len(words) == len(row.text.split())
        first, last = words[0], words[-1]
        assert first[0] == pytest.approx(BOX["x0"], abs=0.05)
        end = last[0] + font.text_length(last[2], fontsize=layout.size)
        assert end == pytest.approx(BOX["x1"], abs=0.05)
        assert (row.x0, row.x1) == (pytest.approx(BOX["x0"]), pytest.approx(BOX["x1"]))


def test_the_last_row_of_each_paragraph_is_not_spread(store):
    _, obj = _styled(store, alignment="JUSTIFY")
    layout = obj.drawn_layout
    ends = [r for r in layout.rows if r.end]
    # The row before her line break, and the block's last row.
    assert len(ends) == 2 and layout.rows[-1].end
    for row in ends:
        (entry,) = [e for e in layout.lines if abs(e[1] - row.baseline) < 0.01]
        assert entry[0] == pytest.approx(BOX["x0"], abs=0.01)
        assert entry[2] == row.text
        assert row.x1 < BOX["x1"] - 1


def test_a_row_of_one_word_is_not_spread(store):
    box = {"x0": 72.0, "y0": 420.0, "x1": 200.0, "y1": 560.0}
    # At 12pt the long word fits the box only alone: it gets a row of its own.
    _, obj = _styled(
        store, text="a Electroencefalografista b", box=box, size=12.0, alignment="JUSTIFY",
    )
    layout = obj.drawn_layout
    assert [r.text for r in layout.rows] == ["a", "Electroencefalografista", "b"]
    assert all(e[0] == pytest.approx(box["x0"], abs=0.01) for e in layout.lines)


def test_each_row_says_whether_it_was_spread(store):
    from ct_engine.api.routes import _block_out

    _, obj = _styled(store, alignment="JUSTIFY")
    rows = obj.drawn_layout.rows
    assert [r.spread for r in rows] == [not r.end for r in rows]
    assert [r["spread"] for r in _block_out(obj)["rows"]] == [r.spread for r in rows]
    # A word wider than the box cannot be spread, though the row has two.
    box = {"x0": 72.0, "y0": 420.0, "x1": 120.0, "y1": 560.0}
    _, narrow = _styled(
        store, text="Electroencefalografista dos\nb", box=box, size=12.0, alignment="JUSTIFY",
    )
    assert not any(r.spread for r in narrow.drawn_layout.rows)


def test_wrap_marks_the_rows_that_end_a_paragraph():
    font = REGISTRY.resolve(SPANISH)
    lines, ends = wrap_with_ends(SPANISH + "\n\nFin.", font, 11, 200)
    assert len(lines) == len(ends)
    assert ends[-1] and lines[-1] == "Fin."
    assert lines[-2] == "" and ends[-2]
    first_paragraph = SPANISH.split("\n")[0].split()[-1]
    i = next(k for k, line in enumerate(lines) if line.endswith(first_paragraph))
    assert ends[i] and not ends[0]


def test_arabic_is_spread_and_its_last_row_sits_on_the_right():
    data, pages = prepared_pages("01_clean_worksheet.pdf", target="ar")
    obj = next(t for t in pages[0].text_objects if t.flat_text.startswith("Directions"))
    obj.translated_text = (
        "اقرأ كل سؤال بعناية واكتب إجابتك على السطور أدناه باستخدام جمل كاملة "
        "ومفردات الوحدة التي تعلمتها في الفصل هذا الأسبوع"
    )
    obj.manual_bbox = Rect(x0=72, y0=420, x1=260, y1=560)
    obj.manual_alignment = TextAlignment.JUSTIFY
    obj.manual_font_size, obj.manual_size_locked = 11.0, True
    rebuild_document(data, pages, target_lang="ar")
    rows = obj.drawn_layout.rows
    assert len(rows) >= 3
    for row in rows[:-1]:
        assert (row.x0, row.x1) == (pytest.approx(72.0), pytest.approx(260.0))
    assert rows[-1].x1 == pytest.approx(260.0, abs=0.01) and rows[-1].x0 > 80


def test_chinese_justified_is_drawn_from_its_start(store):
    text = "光合作用是植物利用阳光把水和二氧化碳变成葡萄糖和氧气的过程。" * 3
    assert not justifiable(text)
    _, obj = _styled(store, text=text, lang="zh-Hans", alignment="JUSTIFY")
    layout = obj.drawn_layout
    assert len(layout.rows) >= 2
    # One entry per row, each from the left edge: nothing spread.
    assert len(layout.lines) == len([r for r in layout.rows if r.text])
    assert all(e[0] == pytest.approx(BOX["x0"], abs=0.01) for e in layout.lines)


def test_a_justified_source_is_still_drawn_from_the_left(store):
    manifest = prepared(store)
    obj = _directions(manifest)
    obj.alignment = TextAlignment.JUSTIFY
    apply_edit(manifest, 1, obj.id, SPANISH, store=store, target_lang="es")
    assert obj.drawn_alignment is TextAlignment.LEFT
    assert len(obj.drawn_layout.lines) == len([r for r in obj.drawn_layout.rows if r.text])


# --- 7. Underline and highlight, in pixels --------------------------------
def _yellow(arr):
    return (arr[..., 0] > 230) & (arr[..., 1] > 230) & (arr[..., 2] < 90)


def test_the_highlight_covers_each_row_and_stays_inside_the_frame(store):
    manifest, obj = _styled(store, highlight="#ffff00")
    arr, _ = _page_pixels(manifest, store)
    yellow = _yellow(arr)
    s = obj.drawn_layout.size
    for row in obj.drawn_layout.rows:
        band = yellow[_px(row.baseline - 0.7 * s):_px(row.baseline + 0.1 * s),
                      _px(row.x0) + 1:_px(row.x1) - 1]
        assert band.mean() > 0.3
        # Past the row's last word: no highlight.
        assert not yellow[_px(row.baseline - 0.4 * s), _px(row.x1) + 3]
    f = obj.drawn_bbox
    cols = slice(_px(f.x0), _px(f.x1))
    assert not yellow[: _px(f.y0) - 1, cols].any()
    assert not yellow[_px(f.y1) + 1:, cols].any()


def test_the_underline_is_a_rule_just_below_each_baseline(store):
    manifest, obj = _styled(store, underline=True)
    dpi = 288
    arr, _ = _page_pixels(manifest, store, dpi=dpi)
    dark = arr.max(axis=2) < 100
    s = obj.drawn_layout.size
    for row in obj.drawn_layout.rows:
        if not row.text:
            continue
        band = dark[_px(row.baseline + 0.08 * s, dpi):_px(row.baseline + 0.14 * s, dpi) + 1,
                    _px(row.x0, dpi) + 2:_px(row.x1, dpi) - 2]
        # Every column across the row has ink in the band: a solid rule.
        assert band.any(axis=0).mean() > 0.98


def test_a_highlight_on_a_filled_box_shows_both(store):
    manifest, obj = _styled(store, highlight="#ffff00", fill="#3366ff")
    arr, _ = _page_pixels(manifest, store)
    row = obj.drawn_layout.rows[0]
    s = obj.drawn_layout.size
    assert _yellow(arr[_px(row.baseline - 0.5 * s), _px(row.x0) + 1:_px(row.x1) - 1]).mean() > 0.3
    below = arr[_px(BOX["y1"]) - 4, _px(BOX["x0"]) + 4]
    assert abs(below[0] - 0x33) < 8 and abs(below[2] - 0xff) < 8


def test_words_on_a_dark_highlight_turn_white(store):
    _, obj = _styled(store, highlight="#1a237e")
    assert obj.drawn_layout.color == (1.0, 1.0, 1.0)
    _, obj = _styled(store, highlight="#1a237e", color="#000000")
    assert obj.drawn_layout.color == (0.0, 0.0, 0.0)  # her colour wins


def test_the_clean_background_has_no_highlight(store):
    manifest, obj = _styled(store, highlight="#ffff00", underline=True)
    plates: dict = {}
    arr, _ = _page_pixels(manifest, store, plates=plates)
    assert _yellow(arr).sum() > 100
    assert _yellow(plates[0]).sum() == 0


def test_plate_plus_sprites_is_the_page_with_every_style():
    data, pages = prepared_pages("01_clean_worksheet.pdf", provider=EchoProvider(lambda t: f"{t} ok"))
    texts = [t for t in pages[0].text_objects if t.translated_text]
    texts[5].manual_underline = True
    texts[5].manual_highlight = "#fff59d"
    texts[5].manual_alignment = TextAlignment.JUSTIFY
    texts[6].manual_bbox = Rect(x0=72, y0=170, x1=400, y1=240)
    texts[6].manual_vertical_align = "middle"
    texts[6].manual_line_spacing = 1.5
    texts[6].manual_highlight = "#b3e5fc"
    texts[7].manual_fill = "#ffe0b2"
    texts[7].manual_underline = True
    plates: dict = {}
    scale = 96 / 72

    def on_plate(index, page):
        plates[index] = Image.open(io.BytesIO(page.get_pixmap(dpi=96, alpha=False).tobytes("png"))).convert("RGBA")

    out, _ = rebuild_document(data, pages, target_lang="es", on_plate=on_plate)
    canvas = plates[0].copy()
    for t in sorted(pages[0].text_objects, key=lambda t: t.z_order):
        if t.drawn_bbox is None:
            continue
        sprite = Image.open(io.BytesIO(render_sprite(t, scale=scale))).convert("RGBA")
        canvas.alpha_composite(
            sprite, (round((t.drawn_bbox.x0 - PAD) * scale), round((t.drawn_bbox.y0 - PAD) * scale)),
        )
    with fitz.open(stream=out, filetype="pdf") as doc:
        full = Image.open(io.BytesIO(doc[0].get_pixmap(dpi=96, alpha=False).tobytes("png"))).convert("RGBA")

    def soft(img):
        return np.asarray(img.convert("L").filter(ImageFilter.GaussianBlur(1)), dtype=np.int16)

    assert (np.abs(soft(canvas) - soft(full)) <= 40).mean() > 0.999


def _inside(quad, x, y):
    """Point in a convex quad given in perimeter order."""
    sign = None
    for (ax, ay), (bx, by) in zip(quad, quad[1:] + quad[:1]):
        cross = (bx - ax) * (y - ay) - (by - ay) * (x - ax)
        if abs(cross) < 1e-9:
            continue
        if sign is None:
            sign = cross > 0
        elif (cross > 0) != sign:
            return False
    return True


def test_a_turned_labels_highlight_follows_the_slant_and_reaches_word_files():
    from ct_engine.pipeline.word_pictures import _labels_in
    from test_tilted_text import _read, _rebuild, _translate

    pdf, page, _, _, (label,) = _read(frame=False, big_letter=False)
    _translate(label)
    label.manual_highlight = "#ff00ff"
    label.manual_underline = True
    final, plate, _, _ = _rebuild(pdf, page)
    magenta = (final[..., 0] > 200) & (final[..., 1] < 80) & (final[..., 2] > 200)
    assert magenta.sum() > 200
    assert not ((plate[..., 0] > 200) & (plate[..., 1] < 80) & (plate[..., 2] > 200)).any()
    # Every highlight and underline pixel lies in the turned frame (allowing
    # for anti-aliasing), so the area Word files copy holds them all.
    frame = grow_quad(label.drawn_layout.quad, 1.0)
    areas = [a for a in _labels_in(page, page.image_objects[0]) if isinstance(a, list)]
    ys, xs = np.nonzero(magenta)
    for x, y in list(zip(xs, ys))[::25]:
        px, py = (x + 0.5) / 2, (y + 0.5) / 2  # 144 dpi
        assert _inside(frame, px, py)
        assert any(_inside(a, px, py) for a in areas)


# --- 8. Office files ------------------------------------------------------
def test_word_export_carries_underline_and_highlight(store):
    from ct_engine.ingest.docx_parser import DocxParser
    from ct_engine.model.document import DocumentManifest
    from ct_engine.model.enums import TextSource
    from ct_engine.pipeline.word_export import build_office_export

    data = corpus_file("20_worksheet.docx")
    store.put_bytes("documents/w/source.docx", data)
    manifest = DocumentManifest(
        document_id="w", filename="w.docx", content_type="x", byte_size=len(data),
        sha256="0", source_key="documents/w/source.docx",
    )
    manifest.office_source_key, manifest.office_format = "documents/w/source.docx", "docx"
    with DocxParser().open_bytes(data) as doc:
        manifest.pages = [doc.parse_page(i) for i in range(doc.page_count)]
    # The page blocks stand in for what the editor showed: each paragraph
    # became one block with the same words.
    for p in manifest.pages:
        for i, t in enumerate(p.text_objects):
            t.source, t.reading_index = TextSource.NATIVE, i
            t.translated_text = f"<{t.flat_text}>"
            t.translation_status = TranslationStatus.TRANSLATED
    target = next(
        t for t in manifest.pages[0].text_objects
        if t.text.strip() == "Cells and Energy: Review Packet"
    )
    target.manual_underline, target.manual_highlight = True, "#fff59d"
    result = build_office_export(manifest, store=store, target_lang="es")
    xml = zipfile.ZipFile(io.BytesIO(result.data)).read("word/document.xml").decode()
    run = xml[xml.index("&lt;Cells and Energy") - 600: xml.index("&lt;Cells and Energy")]
    assert '<w:u w:val="single"/>' in run
    assert 'w:fill="FFF59D"' in run
    # In schema order: u before shd.
    assert run.index("<w:u ") < run.index("<w:shd ")


def test_docx_run_properties_keep_their_schema_order():
    from lxml import etree

    from ct_engine.ingest.docx_walk import W
    from ct_engine.reconstruct.docx_builder import _apply_manual_style

    run = etree.fromstring(
        f'<w:r xmlns:w="{W[1:-1]}"><w:rPr><w:b/><w:lang w:val="en-US"/></w:rPr><w:t>x</w:t></w:r>'
    )
    obj = SimpleNamespace(
        manual_color="#112233", manual_font_size=12.0,
        manual_underline=True, manual_highlight="#fff59d",
    )
    _apply_manual_style(run, obj)
    names = [etree.QName(c).localname for c in run.find(f"{W}rPr")]
    assert names == ["b", "color", "sz", "u", "shd", "lang"]


def test_her_highlight_replaces_a_highlighter_colour_the_run_had():
    from lxml import etree

    from ct_engine.ingest.docx_walk import W
    from ct_engine.reconstruct.docx_builder import _apply_manual_style

    run = etree.fromstring(
        f'<w:r xmlns:w="{W[1:-1]}"><w:rPr><w:highlight w:val="green"/></w:rPr>'
        f"<w:t>x</w:t></w:r>"
    )
    obj = SimpleNamespace(
        manual_color=None, manual_font_size=None,
        manual_underline=False, manual_highlight="#fff59d",
    )
    _apply_manual_style(run, obj)
    rpr = run.find(f"{W}rPr")
    assert rpr.find(f"{W}highlight") is None
    assert rpr.find(f"{W}shd").get(f"{W}fill") == "FFF59D"


def test_powerpoint_export_carries_underline_and_highlight_in_order():
    from pptx import Presentation

    from test_pptx import DECK, _pages, _slide_xml, _translate
    from ct_engine.reconstruct.pptx_builder import rebuild_pptx

    pages = _translate(_pages(), lambda t: t)
    target = next(o for o in pages[0].text_objects if "Photosynthesis" in o.text)
    target.manual_color = "#1155cc"
    target.manual_underline = True
    target.manual_highlight = "#fff59d"
    result = rebuild_pptx(corpus_file(DECK), pages)
    xml = _slide_xml(result.data, 1)
    assert 'u="sng"' in xml
    at = xml.index('<a:highlight><a:srgbClr val="FFF59D"/></a:highlight>')
    rpr = xml[xml.rindex("<a:rPr", 0, at): xml.index("</a:rPr>", at)]
    assert rpr.index("<a:solidFill>") < rpr.index("<a:highlight>")
    for later in ("<a:latin", "<a:ea", "<a:cs"):
        if later in rpr:
            assert rpr.index("<a:highlight>") < rpr.index(later)
    Presentation(io.BytesIO(result.data))  # reopens


def test_editable_powerpoint_notes_only_the_styled_pages():
    from pptx import Presentation

    from ct_engine.reconstruct.layered_pptx import (
        STYLED_TEXT_NOTE,
        export_layered_pptx,
        styled_pages,
    )

    with fitz.open() as pdf:
        for _ in range(2):
            pdf.new_page(width=612, height=792).insert_text((72, 90), "Texto", fontsize=16)
        data = pdf.tobytes()
    deck = Presentation(io.BytesIO(export_layered_pptx(data, style_notes={1})))
    assert not deck.slides[0].has_notes_slide
    assert deck.slides[1].notes_slide.notes_text_frame.text == STYLED_TEXT_NOTE

    drawn = DrawnLayout(size=10.0)

    def block():
        return SimpleNamespace(
            drawn_layout=drawn, suppressed=False, manual_underline=False,
            manual_highlight=None, drawn_alignment=TextAlignment.LEFT,
            translated_text="hola mundo",
        )

    pages = [
        SimpleNamespace(index=0, text_objects=[block()]),
        SimpleNamespace(index=1, text_objects=[block()]),
    ]
    pages[1].text_objects[0].manual_underline = True
    assert styled_pages(pages) == {1}
    pages[0].text_objects[0].drawn_alignment = TextAlignment.JUSTIFY
    assert styled_pages(pages) == {0, 1}
    pages[0].text_objects[0].translated_text = "你好世界"
    assert styled_pages(pages) == {1}
